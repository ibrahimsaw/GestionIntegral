"""Aperçu et exports des périodes de panne et de maintenance."""

import calendar
from datetime import date
import io
from itertools import groupby

from django.contrib import messages
from django.contrib.auth.mixins import LoginRequiredMixin
from django.core.paginator import Paginator
from django.http import HttpResponse
from django.shortcuts import redirect, render
from django.urls import reverse
from django.template.loader import render_to_string
from django.utils import timezone
from django.views import View
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from weasyprint import HTML

from inventory.models import Maintenance, Support


def _add_months(start_date, months):
    month_index = start_date.month - 1 + months
    year = start_date.year + month_index // 12
    month = month_index % 12 + 1
    day = min(start_date.day, calendar.monthrange(year, month)[1])
    return date(year, month, day)


def build_panne_periods(filters):
    date_debut = date.fromisoformat(filters['date_debut']) if filters['date_debut'] else None
    months = filters.get('duree_mois', '').strip()
    if months:
        if not date_debut:
            raise ValueError('Renseignez une date de début pour calculer la durée en mois.')
        try:
            months = int(months)
        except ValueError as error:
            raise ValueError('La durée doit être un nombre entier de mois.') from error
        if months < 1:
            raise ValueError('La durée doit être d’au moins un mois.')
        date_fin = _add_months(date_debut, months)
    else:
        date_fin = date.fromisoformat(filters['date_fin']) if filters['date_fin'] else None
    if date_debut and date_fin and date_debut > date_fin:
        raise ValueError('La date de début doit être antérieure ou égale à la date de fin.')

    maintenances = Maintenance.objects.select_related(
        'support', 'face', 'reparation_commencee_par',
    ).order_by('support_id', 'face_id', 'date_intervention', 'pk')
    if filters['support']:
        maintenances = maintenances.filter(support_id=filters['support'])
    if filters['face']:
        maintenances = maintenances.filter(face__label=filters['face'])

    periods = []

    def group_key(maintenance):
        return maintenance.support_id, maintenance.face_id

    for _, group in groupby(maintenances, key=group_key):
        maintenance_group = list(group)
        support = maintenance_group[0].support
        face = maintenance_group[0].face
        incident_start = None
        incident_maintenance = None
        incident_state = None

        for maintenance in maintenance_group:
            if maintenance.etat_apres in (Support.ETAT_A_MAINTENIR, Support.ETAT_PANNE):
                if incident_start is None:
                    incident_start = maintenance.date_intervention
                    incident_maintenance = maintenance
                    incident_state = maintenance.etat_apres
                else:
                    incident_state = maintenance.etat_apres
            elif maintenance.etat_apres == Support.ETAT_BON and incident_start is not None:
                periods.append({
                    'support': support,
                    'face': face,
                    'debut': incident_start,
                    'fin': maintenance.date_intervention,
                    'duree': maintenance.date_intervention - incident_start,
                    'resolue': True,
                    'maint_debut': incident_maintenance,
                    'etat_debut': incident_maintenance.etat_apres,
                    'etat_incident': incident_state,
                    'maint_bon': maintenance,
                    'reparation_commencee_le': incident_maintenance.reparation_commencee_le,
                    'reparation_commencee_par': incident_maintenance.reparation_commencee_par,
                })
                incident_start = None
                incident_maintenance = None
                incident_state = None

        if incident_start is not None:
            periods.append({
                'support': support,
                'face': face,
                'debut': incident_start,
                'fin': None,
                'duree': timezone.now() - incident_start,
                'resolue': False,
                'maint_debut': incident_maintenance,
                'etat_debut': incident_maintenance.etat_apres,
                'etat_incident': incident_state,
                'maint_bon': None,
                'reparation_commencee_le': incident_maintenance.reparation_commencee_le,
                'reparation_commencee_par': incident_maintenance.reparation_commencee_par,
            })

    if filters['statut'] == 'resolue':
        periods = [period for period in periods if period['resolue']]
    elif filters['statut'] == 'en_cours':
        periods = [period for period in periods if not period['resolue']]

    now = timezone.now()
    periods = [
        period for period in periods
        if (
            not date_debut
            or timezone.localtime(period['fin'] or now).date() >= date_debut
        )
        and (
            not date_fin
            or timezone.localtime(period['debut']).date() <= date_fin
        )
    ]
    periods.sort(key=lambda period: period['debut'], reverse=True)
    periods.sort(key=lambda period: period['etat_incident'] != Support.ETAT_PANNE)
    periods.sort(key=lambda period: period['resolue'])
    return periods, date_debut, date_fin


def _report_filters(request):
    return {
        'support': request.GET.get('support', ''),
        'face': request.GET.get('face', ''),
        'statut': request.GET.get('statut', ''),
        'date_debut': request.GET.get('date_debut', ''),
        'date_fin': request.GET.get('date_fin', ''),
        'duree_mois': request.GET.get('duree_mois', ''),
    }


def _support_filter_label(support_id):
    if not support_id:
        return ''
    support = Support.objects.filter(pk=support_id).values_list('code', 'nom').first()
    return f'{support[0]} — {support[1]}' if support else support_id


class PannesReportView(LoginRequiredMixin, View):
    template_name = 'reports/apercu_pannes.html'

    def get(self, request):
        filters = _report_filters(request)
        try:
            periods, date_debut, date_fin = build_panne_periods(filters)
        except ValueError as error:
            messages.error(request, str(error) or 'Les dates du rapport sont invalides.')
            return redirect('pannes_report_preview')

        if filters['duree_mois']:
            filters['date_fin'] = date_fin.isoformat()
        total = len(periods)
        resolved = sum(period['resolue'] for period in periods)
        repair_started = sum(
            not period['resolue'] and bool(period['reparation_commencee_le'])
            for period in periods
        )
        page_obj = Paginator(periods, 20).get_page(request.GET.get('page'))
        export_params = request.GET.copy()
        export_params.pop('page', None)
        export_params.pop('format', None)
        if filters['duree_mois']:
            export_params['date_fin'] = filters['date_fin']

        return render(request, self.template_name, {
            'title': 'Rapport des pannes et maintenances',
            'periodes': page_obj,
            'page_obj': page_obj,
            'paginator': page_obj.paginator,
            'is_paginated': page_obj.has_other_pages(),
            'nb_total': total,
            'nb_resolues': resolved,
            'nb_en_cours': total - resolved,
            'nb_a_reparer': total - resolved - repair_started,
            'nb_reparation_en_cours': repair_started,
            'supports_liste': Support.objects.filter(
                maintenances__isnull=False,
            ).distinct().order_by('code'),
            'filters': filters,
            'support_label': _support_filter_label(filters['support']),
            'date_debut_obj': date_debut,
            'date_fin_obj': date_fin,
            'date_generation': timezone.localtime(),
            'export_pdf_url': f"{reverse('pannes_report_pdf')}?{export_params.urlencode()}",
            'export_excel_url': f"{reverse('pannes_report_excel')}?{export_params.urlencode()}",
        })


class _PannesReportExportView(LoginRequiredMixin, View):
    def get_periods(self, request):
        try:
            return build_panne_periods(_report_filters(request))
        except ValueError as error:
            messages.error(request, str(error) or 'Les dates du rapport sont invalides.')
            return None


class ExportPannesPdfView(_PannesReportExportView):
    def get(self, request):
        result = self.get_periods(request)
        if result is None:
            return redirect('pannes_report_preview')
        periods, date_debut, date_fin = result
        resolved = sum(period['resolue'] for period in periods)
        repair_started = sum(
            not period['resolue'] and bool(period['reparation_commencee_le'])
            for period in periods
        )
        filters = _report_filters(request)
        html = render_to_string(
            'reports/pannes_report_pdf_export.html',
            {
                'periodes': periods,
                'nb_total': len(periods),
                'nb_resolues': resolved,
                'nb_en_cours': len(periods) - resolved,
                'nb_a_reparer': len(periods) - resolved - repair_started,
                'nb_reparation_en_cours': repair_started,
                'date_debut': date_debut,
                'date_fin': date_fin,
                'date_generation': timezone.localtime(),
                'filters': filters,
                'support_label': _support_filter_label(filters['support']),
            },
            request=request,
        )
        pdf = HTML(
            string=html,
            base_url=request.build_absolute_uri('/'),
        ).write_pdf()
        response = HttpResponse(pdf, content_type='application/pdf')
        response['Content-Disposition'] = 'attachment; filename="rapport_pannes_maintenances.pdf"'
        return response


class ExportPannesExcelView(_PannesReportExportView):
    def get(self, request):
        result = self.get_periods(request)
        if result is None:
            return redirect('pannes_report_preview')
        periods, _, _ = result

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'Pannes et maintenances'
        sheet.append([
            'Support', 'Nom du support', 'Face', 'Type d’incident',
            'Date de début', 'Date de fin / réparation', 'Durée (jours)',
            'Statut', 'Début de la réparation', 'Constat initial', 'Rapport de réparation',
        ])
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill(fill_type='solid', fgColor='932E2B')
            cell.alignment = Alignment(horizontal='center', vertical='center')

        state_labels = dict(Maintenance.ETAT_CHOICES)
        for period in periods:
            sheet.append([
                period['support'].code,
                period['support'].nom,
                f"Face {period['face'].label}" if period['face'] else 'Support complet',
                state_labels.get(period['etat_debut'], period['etat_debut']),
                timezone.localtime(period['debut']).replace(tzinfo=None),
                timezone.localtime(period['fin']).replace(tzinfo=None) if period['fin'] else None,
                round(period['duree'].total_seconds() / 86400, 2),
                (
                    'Résolue' if period['resolue']
                    else 'En cours de réparation' if period['reparation_commencee_le']
                    else 'À réparer'
                ),
                timezone.localtime(period['reparation_commencee_le']).replace(tzinfo=None)
                if period['reparation_commencee_le'] else None,
                period['maint_debut'].description,
                period['maint_bon'].description if period['maint_bon'] else '',
            ])

        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            width = max(len(str(cell.value or '')) for cell in column)
            sheet.column_dimensions[get_column_letter(column[0].column)].width = min(max(width + 2, 12), 45)
        for row in sheet.iter_rows(min_row=2, min_col=5, max_col=6):
            for cell in row:
                cell.number_format = 'dd/mm/yyyy hh:mm'
        for row in sheet.iter_rows(min_row=2, min_col=9, max_col=9):
            row[0].number_format = 'dd/mm/yyyy hh:mm'

        output = io.BytesIO()
        workbook.save(output)
        response = HttpResponse(
            output.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        response['Content-Disposition'] = 'attachment; filename="rapport_pannes_maintenances.xlsx"'
        return response
