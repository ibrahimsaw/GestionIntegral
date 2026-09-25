"""
Rapport global de la liste des marchés — vue, contexte et export PDF/Excel.
"""

from datetime import date, datetime
import io

from django.contrib.auth.mixins import LoginRequiredMixin
from django.core.exceptions import PermissionDenied
from django.http import HttpResponse
from django.shortcuts import render
from django.template.loader import render_to_string
from django.views import View
from django.utils import timezone
from django.db.models import Q
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from weasyprint import HTML

from accounts.decorators import ClientStaffRequiredMixin
from campaigns.models import Campagne
from inventory.models import Marche, Emplacement, Support


def _build_context_marches(filters: dict) -> dict:
    today = date.today()
    now = timezone.now()

    # ── Base queryset ────────────────────────────────────────────
    qs = (
        Marche.objects
        .prefetch_related(
            "emplacements__support_installe__faces",
            "emplacements__lignes_campagne__campagne__client",
        )
        .order_by("ville", "quartier", "nom")
    )

    # ── Filtres ──────────────────────────────────────────────────
    ville = filters.get("ville", "").strip()
    quartier = filters.get("quartier", "").strip()
    statut = filters.get("statut", "").strip()
    occupation = filters.get("occupation", "").strip()
    q = filters.get("q", "").strip()

    if ville:
        qs = qs.filter(ville__iexact=ville)
    if quartier:
        qs = qs.filter(quartier__iexact=quartier)
    if statut:
        if statut.lower() in ("actif", "1", "true"):
            qs = qs.filter(actif=True)
        elif statut.lower() in ("inactif", "0", "false"):
            qs = qs.filter(actif=False)
    if q:
        qs = qs.filter(
            Q(nom__icontains=q)
            | Q(adresse__icontains=q)
            | Q(quartier__icontains=q)
            | Q(ville__icontains=q)
        )

    all_villes = sorted(list(Marche.objects.exclude(ville="").values_list("ville", flat=True).distinct()))
    all_quartiers = sorted(list(Marche.objects.exclude(quartier="").values_list("quartier", flat=True).distinct()))

    # ── Construction des données par marché ──────────────────────
    marches_data = []
    total_emplacements_global = 0
    total_occupes_global = 0
    total_libres_global = 0
    total_marches_actifs = 0

    for m in qs:
        emplacements_list = []
        empls = list(m.emplacements.all())
        nb_empl = len(empls)
        nb_occ = 0

        for emp in empls:
            is_occ = emp.est_dans_campagne
            campagne_active = None
            if is_occ:
                nb_occ += 1
                campagne_active = (
                    Campagne.objects.filter(
                        lignes__emplacement=emp,
                        date_debut__lte=today,
                        date_fin__gte=today,
                    )
                    .select_related("client")
                    .first()
                )

            support_installe = getattr(emp, "support_installe", None)
            fmt_str = "—"
            if support_installe:
                if hasattr(support_installe, 'format_support') and support_installe.format_support:
                    fmt_str = support_installe.format_support.dimensions
                elif hasattr(support_installe, 'format') and support_installe.format:
                    fmt_str = support_installe.format
                elif hasattr(support_installe, 'type_panneau') and support_installe.type_panneau:
                    fmt_str = support_installe.type_panneau

            emplacements_list.append({
                "emplacement": emp,
                "code": emp.code,
                "notes": emp.notes,
                "format": fmt_str,
                "statut": "occupe" if is_occ else "libre",
                "is_occupe": is_occ,
                "campagne": campagne_active,
                "support": support_installe,
            })

        nb_libres = nb_empl - nb_occ
        taux_occ = round((nb_occ / nb_empl * 100), 1) if nb_empl > 0 else 0.0

        # Filtre d'occupation
        if occupation == "occupe" and nb_occ == 0:
            continue
        elif occupation == "libre" and nb_libres == 0:
            continue
        elif occupation == "total_occupe" and (nb_occ != nb_empl or nb_empl == 0):
            continue
        elif occupation == "total_libre" and (nb_libres != nb_empl or nb_empl == 0):
            continue

        if m.actif:
            total_marches_actifs += 1

        total_emplacements_global += nb_empl
        total_occupes_global += nb_occ
        total_libres_global += nb_libres

        marches_data.append({
            "marche": m,
            "pk": m.pk,
            "code": f"MAR-{m.pk:03d}",
            "nom": m.nom,
            "ville": m.ville,
            "quartier": m.quartier,
            "adresse": m.adresse,
            "rayon_metres": m.rayon_metres,
            "latitude": m.latitude,
            "longitude": m.longitude,
            "actif": m.actif,
            "nb_emplacements": nb_empl,
            "nb_occupes": nb_occ,
            "nb_libres": nb_libres,
            "taux_occupation": taux_occ,
            "emplacements": emplacements_list,
        })

    taux_global = (
        round((total_occupes_global / total_emplacements_global * 100), 1)
        if total_emplacements_global > 0
        else 0.0
    )

    return {
        "today": today,
        "marches": marches_data,
        "total_marches": len(marches_data),
        "total_marches_actifs": total_marches_actifs,
        "total_emplacements": total_emplacements_global,
        "total_occupes": total_occupes_global,
        "total_libres": total_libres_global,
        "taux_occupation_global": taux_global,
        "all_villes": all_villes,
        "all_quartiers": all_quartiers,
        "filters": filters,
    }


class MarchesReportView(LoginRequiredMixin, ClientStaffRequiredMixin, View):
    """Aperçu web interactif du rapport global des marchés."""

    def get(self, request, *args, **kwargs):
        filters = {
            "ville": request.GET.get("ville", ""),
            "quartier": request.GET.get("quartier", ""),
            "statut": request.GET.get("statut", ""),
            "occupation": request.GET.get("occupation", ""),
            "q": request.GET.get("q", ""),
        }
        context = _build_context_marches(filters)
        return render(request, "reports/apercu_marches.html", context)


class ExportMarchesPdfView(LoginRequiredMixin, ClientStaffRequiredMixin, View):
    """Génération et téléchargement du rapport PDF global des marchés."""

    def get(self, request, *args, **kwargs):
        filters = {
            "ville": request.GET.get("ville", ""),
            "quartier": request.GET.get("quartier", ""),
            "statut": request.GET.get("statut", ""),
            "occupation": request.GET.get("occupation", ""),
            "q": request.GET.get("q", ""),
        }
        context = _build_context_marches(filters)
        html_string = render_to_string("reports/marches_report_pdf_export.html", context, request=request)
        
        pdf_file = HTML(string=html_string, base_url=request.build_absolute_uri("/")).write_pdf()
        filename = f"rapport_marches_{date.today().strftime('%Y%m%d')}.pdf"
        response = HttpResponse(pdf_file, content_type="application/pdf")
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        return response


class ExportMarchesExcelView(LoginRequiredMixin, ClientStaffRequiredMixin, View):
    """Export Excel de l'état des marchés et leurs emplacements."""

    def get(self, request, *args, **kwargs):
        import openpyxl
        filters = {
            "ville": request.GET.get("ville", ""),
            "quartier": request.GET.get("quartier", ""),
            "statut": request.GET.get("statut", ""),
            "occupation": request.GET.get("occupation", ""),
            "q": request.GET.get("q", ""),
        }
        context = _build_context_marches(filters)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Marchés & Emplacements"

        # Styles
        font_header = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        fill_header = PatternFill(start_color="932E2B", end_color="932E2B", fill_type="solid")

        headers = [
            "Code Marché", "Nom du Marché", "Ville", "Quartier", "Statut",
            "Code Emplacement", "Format", "État Occupation", "Campagne Active"
        ]
        ws.append(headers)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = font_header
            cell.fill = fill_header
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for m in context["marches"]:
            if not m["emplacements"]:
                ws.append([
                    m["code"], m["nom"], m["ville"], m["quartier"], "Actif" if m["actif"] else "Inactif",
                    "—", "—", "Aucun emplacement", "—"
                ])
            else:
                for emp in m["emplacements"]:
                    camp_nom = emp["campagne"].nom if emp["campagne"] else "—"
                    ws.append([
                        m["code"], m["nom"], m["ville"], m["quartier"], "Actif" if m["actif"] else "Inactif",
                        emp["code"], emp["format"], "Occupé" if emp["is_occupe"] else "Libre", camp_nom
                    ])

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        filename = f"rapport_marches_{date.today().strftime('%Y%m%d')}.xlsx"
        response = HttpResponse(
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        return response

