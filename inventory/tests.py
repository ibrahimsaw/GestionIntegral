import datetime
import io
import zipfile
from importlib import import_module
import openpyxl
from django.apps import apps
from django.contrib.auth.models import AnonymousUser
from django.contrib.auth import get_user_model
from django.contrib.auth.models import AnonymousUser
from django.contrib.messages import get_messages
from django.core.cache import cache
from django.template.loader import render_to_string
from django.test import RequestFactory, TestCase
from django.urls import reverse
from django.utils import timezone
from unittest.mock import patch

from campaigns.forms import LigneCampagneForm
from campaigns.models import Campagne, Client, LigneCampagne
from accounts.models import AuditLog
from inventory.forms import EcranNumeriqueForm, MaintenanceForm, SupportForm
from inventory.models import (
    Emplacement,
    EcranNumerique,
    FacePanneau,
    FormatSupport,
    Marche,
    Maintenance,
    ObjectifSpotsEcranAnnuel,
    Support,
    SupportDeletionRequest,
)
from inventory.views import (
    FormatSupportDetailView,
    FormatSupportListView,
    SupportListView,
    SupportDetailView,
)
from planning.models import LogDiffusion
from campaigns.context_processors import alerts_data
from geoad.context_processors import inventory_villes


class EcranAnnualObjectiveFormTests(TestCase):
    def test_technical_form_has_no_annual_goal_field(self):
        screen_format = FormatSupport.objects.create(
            code='SCR-CELLS',
            dimensions='6x4',
            categorie='Ecran',
        )
        support = Support.objects.create(
            code='SCREEN-ANNUAL-GOAL',
            nom='Écran avec objectif annuel',
            type_support=Support.TYPE_ECRAN,
            latitude=12.34,
            longitude=-1.52,
        )
        screen = EcranNumerique.objects.create(support=support)
        old_goal = ObjectifSpotsEcranAnnuel.objects.create(
            ecran=screen,
            annee=timezone.localdate().year - 1,
            nombre_spots=24000,
        )
        form = EcranNumeriqueForm(
            data={
                'type_ecran': 'exterieur',
                'resolution': 'fullhd',
                'cellule': screen_format.code,
                'heure_allumage': '06:00',
                'heure_extinction': '23:00',
            },
            instance=screen,
        )

        self.assertTrue(form.is_valid(), form.errors)
        self.assertNotIn('objectif_spots_annuel', form.fields)
        form.save()
        old_goal.refresh_from_db()
        self.assertEqual(old_goal.nombre_spots, 24000)
        self.assertFalse(
            screen.objectifs_annuels.filter(annee=timezone.localdate().year).exists()
        )


class EcranAnnualObjectiveManagementTests(TestCase):
    def setUp(self):
        self.support = Support.objects.create(
            code='SCREEN-OBJECTIVES',
            nom='Écran objectifs annuels',
            type_support=Support.TYPE_ECRAN,
            latitude=12.34,
            longitude=-1.52,
        )
        self.screen = EcranNumerique.objects.create(support=self.support)
        self.url = reverse('ecran_objectifs_annuels', args=[self.support.uuid])
        self.staff = get_user_model().objects.create_user(
            username='screen-objectives-staff',
            password='test-password',
            role='staff',
        )

    def test_page_defaults_new_targets_and_lists_history(self):
        previous = ObjectifSpotsEcranAnnuel.objects.create(
            ecran=self.screen,
            annee=timezone.localdate().year - 1,
            nombre_spots=24000,
        )
        self.client.force_login(self.staff)

        response = self.client.get(self.url)

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'value="200000"')
        self.assertContains(response, str(previous.annee))
        self.assertContains(response, '24000')

    def test_saving_a_year_updates_only_that_year_and_preserves_history(self):
        previous = ObjectifSpotsEcranAnnuel.objects.create(
            ecran=self.screen,
            annee=timezone.localdate().year - 1,
            nombre_spots=24000,
        )
        self.client.force_login(self.staff)

        response = self.client.post(self.url, {
            'annee': previous.annee,
            'nombre_spots': 36500,
        })

        self.assertRedirects(
            response,
            f'{self.url}?annee={previous.annee}',
        )
        previous.refresh_from_db()
        self.assertEqual(previous.nombre_spots, 36500)
        self.assertEqual(
            ObjectifSpotsEcranAnnuel.objects.filter(ecran=self.screen).count(),
            1,
        )
        next_year = timezone.localdate().year
        self.client.post(self.url, {'annee': next_year, 'nombre_spots': 200000})
        self.assertEqual(
            set(self.screen.objectifs_annuels.values_list('annee', flat=True)),
            {previous.annee, next_year},
        )

    def test_non_staff_cannot_manage_annual_targets(self):
        technician = get_user_model().objects.create_user(
            username='screen-objectives-technician',
            password='test-password',
            role='technicien',
        )
        self.client.force_login(technician)

        response = self.client.get(self.url)

        self.assertEqual(response.status_code, 302)
        self.assertFalse(ObjectifSpotsEcranAnnuel.objects.exists())

    def test_screen_detail_shows_dashboard_style_current_and_annual_rates(self):
        today = timezone.localdate()
        ObjectifSpotsEcranAnnuel.objects.create(
            ecran=self.screen,
            annee=today.year,
            nombre_spots=200,
        )
        LogDiffusion.objects.create(
            ecran=self.screen,
            timestamp=timezone.now(),
        )
        self.client.force_login(self.staff)

        response = self.client.get(reverse('support_detail', args=[self.support.uuid]))

        self.assertEqual(response.status_code, 200)
        elapsed_days = (today - today.replace(month=1, day=1)).days + 1
        expected_current_rate = round(1 / (200 * elapsed_days / 365) * 100, 1)
        self.assertIn(
            ('Taux actuel', 'bi-speedometer2', f'{expected_current_rate}%', 'var(--color-primary)'),
            response.context['ecran_stats'],
        )
        self.assertIn(
            ('Taux annuel', 'bi-graph-up', '0.5%', 'var(--color-primary)'),
            response.context['ecran_stats'],
        )
        self.assertContains(response, 'Objectifs annuels')


class PanelOccupancyRateTests(TestCase):
    def test_panel_detail_shows_current_and_annual_occupancy_rates(self):
        today = timezone.localdate()
        support = Support.objects.create(
            code='PANEL-RATES',
            nom='Panneau avec taux',
            type_support=Support.TYPE_PANNEAU,
            latitude=12.34,
            longitude=-1.52,
        )
        face_a = FacePanneau.objects.create(support=support, label='A')
        face_b = FacePanneau.objects.create(support=support, label='B')
        client = Client.objects.create(
            nom='Client taux panneau',
            telephone='00000000',
            email='panel-rates@example.com',
        )
        completed_campaign = Campagne.objects.create(
            client=client,
            nom='Campagne terminée panneau',
            date_debut=today.replace(month=1, day=1),
            date_fin=today.replace(month=1, day=2),
            statut=Campagne.STATUT_TERMINEE,
            type_support='12m2',
        )
        future_campaign = Campagne.objects.create(
            client=client,
            nom='Campagne future panneau',
            date_debut=today + datetime.timedelta(days=1),
            date_fin=today + datetime.timedelta(days=2),
            statut=Campagne.STATUT_A_VENIR,
            type_support='12m2',
        )
        LigneCampagne.objects.create(
            campagne=completed_campaign,
            support=support,
            face=face_a,
        )
        LigneCampagne.objects.create(
            campagne=future_campaign,
            support=support,
            face=face_b,
        )
        staff = get_user_model().objects.create_user(
            username='panel-rates-staff',
            password='test-password',
            role='staff',
        )
        self.client.force_login(staff)

        response = self.client.get(reverse('support_detail', args=[support.uuid]))

        self.assertEqual(response.status_code, 200)
        year_days = (
            today.replace(month=12, day=31)
            - today.replace(month=1, day=1)
        ).days + 1
        elapsed_days = (today - today.replace(month=1, day=1)).days + 1
        self.assertEqual(
            response.context['panneau_stats'],
            [
                (
                    'Taux actuel',
                    'bi-speedometer2',
                    f'{round(2 / (2 * elapsed_days) * 100, 1)}%',
                    'var(--color-primary)',
                ),
                (
                    f'Taux annuel {today.year}',
                    'bi-graph-up',
                    f'{round(4 / (2 * year_days) * 100, 1)}%',
                    'var(--color-primary)',
                ),
            ],
        )
        face_stats = {
            face.label: (face.taux_actuel, face.taux_annuel)
            for face in response.context['faces']
        }
        self.assertEqual(
            face_stats['A'],
            (
                round(2 / elapsed_days * 100, 1),
                round(2 / year_days * 100, 1),
            ),
        )
        self.assertEqual(face_stats['B'], (0.0, round(2 / year_days * 100, 1)))
        self.assertContains(response, "Taux d'occupation")
        self.assertContains(response, f'Taux annuel {today.year}')
        self.assertContains(response, 'Taux actuel :</strong>')


class RepairStartTrackingTests(TestCase):
    def test_technician_cannot_manage_another_users_completed_maintenance(self):
        technician = get_user_model().objects.create_user(
            username='maintenance-read-only-tech',
            password='test-password',
            role='technicien',
        )
        owner = get_user_model().objects.create_user(
            username='maintenance-owner',
            password='test-password',
            role='staff',
        )
        support = Support.objects.create(
            code='MAINTENANCE-READ-ONLY',
            nom='Support consultation seule',
            type_support=Support.TYPE_ECRAN,
            latitude=12.34,
            longitude=-1.52,
        )
        maintenance = Maintenance.objects.create(
            support=support,
            effectue_par=owner,
            etat_apres=Support.ETAT_BON,
        )
        self.client.force_login(technician)

        detail_url = reverse('maintenance_detail', args=[maintenance.pk])
        list_response = self.client.get(reverse('maintenance_list'))
        response = self.client.get(detail_url)

        self.assertContains(list_response, maintenance.support.code)
        self.assertNotContains(
            list_response,
            reverse('maintenance_update', args=[maintenance.pk]),
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, maintenance.support.code)
        self.assertNotContains(response, reverse('maintenance_update', args=[maintenance.pk]))
        self.assertNotContains(
            response,
            reverse('maintenance_quick_repair', args=[maintenance.pk]),
        )
        self.assertEqual(
            self.client.get(reverse('maintenance_update', args=[maintenance.pk])).status_code,
            404,
        )
        self.assertEqual(
            self.client.post(
                reverse('maintenance_quick_repair', args=[maintenance.pk]),
                {'etat_apres': Support.ETAT_BON},
            ).status_code,
            404,
        )
        self.assertEqual(
            self.client.post(
                reverse('maintenance_start_repair', args=[maintenance.pk]),
            ).status_code,
            404,
        )
        maintenance.refresh_from_db()
        self.assertEqual(maintenance.etat_apres, Support.ETAT_BON)
        self.assertEqual(Maintenance.objects.filter(support=support).count(), 1)

    def test_maintenance_creation_is_audited_with_actor_and_values(self):
        technician = get_user_model().objects.create_user(
            username='maintenance-audit-create',
            password='test-password',
            role='technicien',
        )
        support = Support.objects.create(
            code='MAINTENANCE-AUDIT-CREATE',
            nom='Support audit création',
            type_support=Support.TYPE_ECRAN,
            latitude=12.34,
            longitude=-1.52,
        )
        self.client.force_login(technician)

        response = self.client.post(
            reverse('maintenance_create'),
            {
                'support': support.pk,
                'date_intervention': '2026-10-24T10:00',
                'etat_apres': Support.ETAT_PANNE,
                'description': 'Panne détectée sur site',
            },
        )

        self.assertEqual(response.status_code, 302)
        maintenance = Maintenance.objects.get(support=support)
        audit = AuditLog.objects.get(
            user=technician,
            action=AuditLog.ACTION_CREATE,
            object_id=maintenance.pk,
        )
        self.assertIn(support.code, audit.detail)
        self.assertEqual(audit.changes['etat_apres']['apres'], Support.ETAT_PANNE)
        self.assertEqual(audit.changes['description']['apres'], 'Panne détectée sur site')

    def test_maintenance_update_is_audited_with_changed_fields(self):
        technician = get_user_model().objects.create_user(
            username='maintenance-audit-update',
            password='test-password',
            role='technicien',
        )
        support = Support.objects.create(
            code='MAINTENANCE-AUDIT-UPDATE',
            nom='Support audit modification',
            type_support=Support.TYPE_ECRAN,
            latitude=12.34,
            longitude=-1.52,
        )
        maintenance = Maintenance.objects.create(
            support=support,
            effectue_par=technician,
            etat_apres=Support.ETAT_PANNE,
            description='Ancienne description',
        )
        self.client.force_login(technician)

        response = self.client.post(
            reverse('maintenance_update', args=[maintenance.pk]),
            {
                'support': support.pk,
                'date_intervention': maintenance.date_intervention.strftime('%Y-%m-%dT%H:%M'),
                'etat_apres': Support.ETAT_BON,
                'description': 'Panne réparée',
            },
        )

        self.assertEqual(response.status_code, 302)
        audit = AuditLog.objects.get(
            user=technician,
            action=AuditLog.ACTION_ETAT_CHANGE,
            object_id=maintenance.pk,
        )
        self.assertEqual(
            audit.changes['etat_apres'],
            {'avant': Support.ETAT_PANNE, 'apres': Support.ETAT_BON},
        )
        self.assertEqual(
            audit.changes['description'],
            {'avant': 'Ancienne description', 'apres': 'Panne réparée'},
        )

    def test_quick_repair_and_repair_start_are_audited(self):
        technician = get_user_model().objects.create_user(
            username='maintenance-audit-repair',
            password='test-password',
            role='technicien',
        )
        support = Support.objects.create(
            code='MAINTENANCE-AUDIT-REPAIR',
            nom='Support audit réparation',
            type_support=Support.TYPE_ECRAN,
            latitude=12.34,
            longitude=-1.52,
        )
        maintenance = Maintenance.objects.create(
            support=support,
            effectue_par=technician,
            etat_apres=Support.ETAT_PANNE,
        )
        self.client.force_login(technician)

        response = self.client.post(
            reverse('maintenance_quick_repair', args=[maintenance.pk]),
            {
                'etat_apres': Support.ETAT_BON,
                'description': 'Écran remis en service',
            },
        )

        self.assertEqual(response.status_code, 302)
        progress = Maintenance.objects.exclude(pk=maintenance.pk).get(support=support)
        progress_audit = AuditLog.objects.get(
            user=technician,
            action=AuditLog.ACTION_ETAT_CHANGE,
            object_id=progress.pk,
        )
        self.assertEqual(
            progress_audit.changes['etat_element'],
            {'avant': Support.ETAT_PANNE, 'apres': Support.ETAT_BON},
        )

        second_support = Support.objects.create(
            code='MAINTENANCE-AUDIT-START',
            nom='Support audit démarrage',
            type_support=Support.TYPE_ECRAN,
            latitude=12.35,
            longitude=-1.53,
        )
        open_maintenance = Maintenance.objects.create(
            support=second_support,
            effectue_par=technician,
            etat_apres=Support.ETAT_PANNE,
        )
        response = self.client.post(
            reverse('maintenance_start_repair', args=[open_maintenance.pk]),
        )

        self.assertEqual(response.status_code, 302)
        start_audit = AuditLog.objects.get(
            user=technician,
            action=AuditLog.ACTION_UPDATE,
            object_id=open_maintenance.pk,
        )
        self.assertIn('Démarrage de la réparation', start_audit.detail)
        self.assertEqual(
            start_audit.changes['reparation_commencee_par']['apres'],
            str(technician),
        )

    def test_client_cannot_manage_maintenance(self):
        client_user = get_user_model().objects.create_user(
            username='maintenance-read-only-client',
            password='test-password',
            role='client',
        )
        owner = get_user_model().objects.create_user(
            username='maintenance-client-test-owner',
            password='test-password',
            role='staff',
        )
        support = Support.objects.create(
            code='MAINTENANCE-CLIENT-READ-ONLY',
            nom='Support client lecture seule',
            type_support=Support.TYPE_ECRAN,
            latitude=12.34,
            longitude=-1.52,
        )
        maintenance = Maintenance.objects.create(
            support=support,
            effectue_par=owner,
            etat_apres=Support.ETAT_PANNE,
        )
        self.client.force_login(client_user)

        list_response = self.client.get(reverse('maintenance_list'))
        detail_response = self.client.get(
            reverse('maintenance_detail', args=[maintenance.pk]),
        )

        self.assertNotContains(
            list_response,
            reverse('maintenance_update', args=[maintenance.pk]),
        )
        self.assertEqual(detail_response.status_code, 200)
        self.assertNotContains(
            detail_response,
            reverse('maintenance_update', args=[maintenance.pk]),
        )
        for route_name in (
            'maintenance_update',
            'maintenance_quick_repair',
            'maintenance_start_repair',
        ):
            with self.subTest(route=route_name):
                url = reverse(route_name, args=[maintenance.pk])
                response = (
                    self.client.get(url)
                    if route_name == 'maintenance_update'
                    else self.client.post(url, {'etat_apres': Support.ETAT_BON})
                )
                self.assertEqual(response.status_code, 404)

        self.assertEqual(Maintenance.objects.filter(support=support).count(), 1)

    def test_technician_can_start_an_open_repair_without_changing_support_state(self):
        technician = get_user_model().objects.create_user(
            username='repair-start-tech',
            password='test-password',
            role='technicien',
        )
        support = Support.objects.create(
            code='REPAIR-START',
            nom='Support réparation',
            type_support=Support.TYPE_ECRAN,
            latitude=12.34,
            longitude=-1.52,
        )
        maintenance = Maintenance.objects.create(
            support=support,
            effectue_par=technician,
            etat_apres=Support.ETAT_PANNE,
        )
        self.client.force_login(technician)

        response = self.client.post(
            reverse('maintenance_start_repair', args=[maintenance.pk]),
        )

        self.assertRedirects(
            response,
            reverse('maintenance_detail', args=[maintenance.pk]),
        )
        maintenance.refresh_from_db()
        support.refresh_from_db()
        self.assertIsNotNone(maintenance.reparation_commencee_le)
        self.assertEqual(maintenance.reparation_commencee_par, technician)
        self.assertEqual(support.etat, Support.ETAT_PANNE)

        detail = self.client.get(reverse('maintenance_detail', args=[maintenance.pk]))
        self.assertContains(detail, 'En cours de réparation')
        self.assertNotContains(detail, 'Démarrer la réparation')

    def test_technician_can_take_over_another_users_started_repair(self):
        technician = get_user_model().objects.create_user(
            username='repair-takeover-tech',
            password='test-password',
            role='technicien',
        )
        owner = get_user_model().objects.create_user(
            username='repair-takeover-owner',
            password='test-password',
            role='staff',
        )
        support = Support.objects.create(
            code='REPAIR-TAKEOVER',
            nom='Support reprise réparation',
            type_support=Support.TYPE_ECRAN,
            latitude=12.35,
            longitude=-1.53,
        )
        maintenance = Maintenance.objects.create(
            support=support,
            effectue_par=owner,
            etat_apres=Support.ETAT_PANNE,
            reparation_commencee_le=datetime.datetime(
                2026, 10, 28, 10, 0, tzinfo=datetime.timezone.utc,
            ),
            reparation_commencee_par=owner,
            description='Diagnostic initial',
        )
        self.client.force_login(technician)

        update_url = reverse('maintenance_update', args=[maintenance.pk])
        list_response = self.client.get(reverse('maintenance_list'))
        self.assertContains(list_response, update_url)
        self.assertEqual(self.client.get(update_url).status_code, 200)
        response = self.client.post(
            update_url,
            {
                'support': support.pk,
                'date_intervention': maintenance.date_intervention.strftime('%Y-%m-%dT%H:%M'),
                'etat_apres': Support.ETAT_PANNE,
                'description': 'Diagnostic repris et complété',
            },
        )
        self.assertEqual(response.status_code, 302)
        maintenance.refresh_from_db()
        self.assertEqual(maintenance.description, 'Diagnostic repris et complété')
        self.assertEqual(maintenance.reparation_commencee_par, owner)

        response = self.client.post(
            reverse('maintenance_quick_repair', args=[maintenance.pk]),
            {
                'etat_apres': Support.ETAT_BON,
                'description': 'Réparation terminée par le technicien repreneur',
            },
        )
        self.assertEqual(response.status_code, 302)
        progress = Maintenance.objects.exclude(pk=maintenance.pk).get(support=support)
        self.assertEqual(progress.effectue_par, technician)
        self.assertEqual(progress.etat_apres, Support.ETAT_BON)
        self.assertEqual(
            progress.description,
            'Réparation terminée par le technicien repreneur',
        )

    def test_technician_can_start_an_unassigned_open_repair(self):
        technician = get_user_model().objects.create_user(
            username='repair-start-unassigned-tech',
            password='test-password',
            role='technicien',
        )
        support = Support.objects.create(
            code='REPAIR-START-UNASSIGNED',
            nom='Support réparation non attribué',
            type_support=Support.TYPE_ECRAN,
            latitude=12.35,
            longitude=-1.53,
        )
        maintenance = Maintenance.objects.create(
            support=support,
            etat_apres=Support.ETAT_PANNE,
        )
        self.client.force_login(technician)

        response = self.client.post(
            reverse('maintenance_start_repair', args=[maintenance.pk]),
        )

        self.assertEqual(response.status_code, 302)
        maintenance.refresh_from_db()
        self.assertIsNotNone(maintenance.reparation_commencee_le)
        self.assertEqual(maintenance.reparation_commencee_par, technician)

    def test_staff_can_start_a_repair(self):
        staff_user = get_user_model().objects.create_user(
            username='repair-start-staff',
            password='test-password',
            role='staff',
        )
        support = Support.objects.create(
            code='REPAIR-START-STAFF',
            nom='Support réparation staff',
            type_support=Support.TYPE_ECRAN,
            latitude=12.36,
            longitude=-1.54,
        )
        maintenance = Maintenance.objects.create(
            support=support,
            etat_apres=Support.ETAT_A_MAINTENIR,
        )
        self.client.force_login(staff_user)

        response = self.client.post(
            reverse('maintenance_start_repair', args=[maintenance.pk]),
        )

        self.assertEqual(response.status_code, 302)
        maintenance.refresh_from_db()
        self.assertIsNotNone(maintenance.reparation_commencee_le)
        self.assertEqual(maintenance.reparation_commencee_par, staff_user)


class PreventiveMaintenanceStateTests(TestCase):
    def test_new_maintenance_applies_only_to_checked_panel_faces(self):
        panel = Support.objects.create(
            code='MAINT-SELECTED-FACES',
            nom='Panneau avec faces sélectionnées',
            type_support=Support.TYPE_PANNEAU,
            latitude=12.34,
            longitude=-1.52,
        )
        face_a = FacePanneau.objects.create(support=panel, label='A')
        face_b = FacePanneau.objects.create(support=panel, label='B')
        user = get_user_model().objects.create_user(
            username='maintenance-admin',
            password='test-password',
            role='admin',
        )
        self.client.force_login(user)

        response = self.client.get(
            reverse('maintenance_create'),
            {'support': panel.pk},
        )
        self.assertContains(response, 'Face A')
        self.assertContains(response, 'Face B')
        self.assertContains(response, 'Ne pas modifier')
        for label in (
            'Bon état',
            'Signal',
            'En panne',
        ):
            self.assertContains(response, label)
        self.assertNotContains(response, 'Fonctionnel — maintenance à prévoir')
        self.assertContains(response, 'type="radio"')
        self.assertContains(response, 'btn-outline-success')
        self.assertContains(response, 'btn-outline-warning')
        self.assertContains(response, 'btn-outline-danger')
        self.assertNotContains(response, '<select name="etat_face_')
        self.assertEqual(response.context['form'][f'etat_face_{face_a.pk}'].value(), '')
        self.assertEqual(response.context['form'][f'etat_face_{face_b.pk}'].value(), '')

        response = self.client.post(
            reverse('maintenance_create'),
            {
                'support': panel.pk,
                'effectue_par': user.pk,
                'date_intervention': '2026-10-24T10:00',
                f'etat_face_{face_a.pk}': '',
                f'etat_face_{face_b.pk}': '',
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(Maintenance.objects.filter(support=panel).count(), 0)

        response = self.client.post(
            reverse('maintenance_create'),
            {
                'support': panel.pk,
                'effectue_par': user.pk,
                'date_intervention': '2026-10-24T10:00',
                f'etat_face_{face_a.pk}': Support.ETAT_PANNE,
                f'etat_face_{face_b.pk}': '',
            },
        )

        self.assertEqual(
            response.status_code,
            302,
            response.context['form'].errors if response.context else response.content,
        )
        self.assertEqual(
            list(
                Maintenance.objects.filter(support=panel).values_list(
                    'face_id', flat=True,
                )
            ),
            [face_a.pk],
        )
        maintenances = Maintenance.objects.filter(support=panel)
        self.assertEqual(
            maintenances.get(face=face_a).etat_apres,
            Support.ETAT_PANNE,
        )
        face_a.refresh_from_db()
        face_b.refresh_from_db()
        self.assertEqual(face_a.etat, Support.ETAT_PANNE)
        self.assertEqual(face_b.etat, Support.ETAT_BON)

        self.client.post(
            reverse(
                'maintenance_quick_repair',
                args=[maintenances.get(face=face_a).pk],
            ),
        )

        response = self.client.post(
            reverse('maintenance_create', args=[panel.pk]),
            {
                'support': panel.pk,
                'effectue_par': user.pk,
                'date_intervention': '2026-10-25T10:00',
                f'etat_face_{face_a.pk}': '',
                f'etat_face_{face_b.pk}': Support.ETAT_A_MAINTENIR,
            },
        )
        self.assertEqual(response.status_code, 302)
        self.assertEqual(Maintenance.objects.filter(support=panel).count(), 3)
        self.assertEqual(
            Maintenance.objects.filter(face=face_a).order_by('-date_intervention').first().etat_apres,
            Support.ETAT_BON,
        )
        self.assertEqual(
            Maintenance.objects.get(face=face_b).etat_apres,
            Support.ETAT_A_MAINTENIR,
        )
        face_b.refresh_from_db()
        self.assertEqual(face_b.etat, Support.ETAT_A_MAINTENIR)

        response = self.client.get(
            reverse(
                'maintenance_detail',
                args=[Maintenance.objects.filter(support=panel).first().pk],
            )
        )
        self.assertContains(
            response,
            f'href="{reverse("maintenance_create", args=[panel.pk])}"',
        )

        maintenance = Maintenance.objects.get(
            support=panel,
            face=face_a,
            etat_apres=Support.ETAT_PANNE,
        )
        update_url = reverse('maintenance_update', args=[maintenance.pk])
        response = self.client.get(update_url)
        self.assertContains(response, 'Face A')
        self.assertNotContains(response, 'name="face"')

        response = self.client.post(
            update_url,
            {
                'support': panel.pk,
                'effectue_par': user.pk,
                'date_intervention': maintenance.date_intervention.strftime('%Y-%m-%dT%H:%M'),
                'etat_apres': Support.ETAT_BON,
                'face': face_b.pk,
            },
        )
        self.assertEqual(response.status_code, 302)
        maintenance.refresh_from_db()
        self.assertEqual(maintenance.face_id, face_a.pk)
        self.assertEqual(maintenance.etat_apres, Support.ETAT_BON)

        response = self.client.get(reverse('maintenance_list'))
        self.assertContains(
            response,
            f'href="{reverse("maintenance_create", args=[panel.pk])}"',
        )

    def test_maintenance_form_offers_only_the_three_intervention_states(self):
        self.assertEqual(
            dict(Support.ETAT_CHOICES)[Support.ETAT_A_MAINTENIR],
            'Signal',
        )
        self.assertNotIn('ETAT_MAINTENANCE', Support.__dict__)
        form = MaintenanceForm()
        self.assertEqual(
            list(form.fields['etat_apres'].choices),
            list(MaintenanceForm.ETAT_CHOICES),
        )
        self.assertNotIn(
            'maintenance',
            [state for state, _ in form.fields['etat_apres'].choices],
        )
        self.assertIn(
            (Support.ETAT_A_MAINTENIR, 'Signal'),
            form.fields['etat_apres'].choices,
        )

    def test_maintenance_marks_screen_functional_but_due_for_maintenance(self):
        screen = Support.objects.create(
            code='MAINT-SCREEN',
            nom='Écran à entretenir',
            type_support=Support.TYPE_ECRAN,
            latitude=12.34,
            longitude=-1.52,
        )

        Maintenance.objects.create(
            support=screen,
            etat_apres=Support.ETAT_A_MAINTENIR,
        )

        screen.refresh_from_db()
        self.assertEqual(screen.etat, Support.ETAT_A_MAINTENIR)

    def test_signal_maintenance_detail_offers_quick_repair(self):
        user = get_user_model().objects.create_user(
            username='maintenance-signal-repair',
            password='test-password',
            role='admin',
        )
        self.client.force_login(user)
        screen = Support.objects.create(
            code='MAINT-SIGNAL-REPAIR',
            nom='Écran avec signal',
            type_support=Support.TYPE_ECRAN,
            latitude=12.34,
            longitude=-1.52,
        )
        maintenance = Maintenance.objects.create(
            support=screen,
            etat_apres=Support.ETAT_A_MAINTENIR,
        )

        response = self.client.get(
            reverse('maintenance_detail', args=[maintenance.pk]),
        )

        self.assertContains(response, 'Ce support a été enregistré comme étant')
        self.assertContains(response, 'SIGNAL')
        self.assertContains(response, 'Avancement !')
        self.assertContains(response, 'Enregistrer l\'avancement')
        self.assertContains(response, 'name="etat_apres"')
        self.assertContains(response, 'État actuel :')
        self.assertContains(response, '<div class="fw-bold text-dark">Signal</div>')
        self.assertNotContains(response, '<option value="a_maintenir"')
        self.assertContains(
            response,
            f'action="{reverse("maintenance_quick_repair", args=[maintenance.pk])}"',
        )

        response = self.client.post(
            reverse('maintenance_quick_repair', args=[maintenance.pk]),
            {'etat_apres': Support.ETAT_A_MAINTENIR},
        )
        self.assertEqual(response.status_code, 302)
        self.assertEqual(Maintenance.objects.filter(support=screen).count(), 1)
        screen.refresh_from_db()
        self.assertEqual(screen.etat, Support.ETAT_A_MAINTENIR)

        response = self.client.post(
            reverse('maintenance_quick_repair', args=[maintenance.pk]),
            {'etat_apres': Support.ETAT_BON},
        )

        self.assertEqual(response.status_code, 302)
        repair = Maintenance.objects.exclude(pk=maintenance.pk).get(
            support=screen,
        )
        self.assertEqual(repair.etat_apres, Support.ETAT_BON)
        screen.refresh_from_db()
        self.assertEqual(screen.etat, Support.ETAT_BON)

    def test_new_panel_maintenance_only_allows_unrepaired_faces(self):
        user = get_user_model().objects.create_user(
            username='maintenance-face-lock',
            password='test-password',
            role='admin',
        )
        self.client.force_login(user)
        panel = Support.objects.create(
            code='MAINT-FACE-LOCK',
            nom='Panneau avec une face signalée',
            type_support=Support.TYPE_PANNEAU,
            latitude=12.34,
            longitude=-1.52,
        )
        signal_face = FacePanneau.objects.create(support=panel, label='A')
        available_face = FacePanneau.objects.create(support=panel, label='B')
        signal_maintenance = Maintenance.objects.create(
            support=panel,
            face=signal_face,
            etat_apres=Support.ETAT_A_MAINTENIR,
        )

        faces_response = self.client.get(
            reverse('api_faces_support', args=[panel.pk]),
        )
        self.assertEqual(faces_response.status_code, 200)
        face_data = next(
            face for face in faces_response.json()['faces']
            if face['id'] == signal_face.pk
        )
        self.assertEqual(face_data['repair_id'], signal_maintenance.pk)

        response = self.client.get(
            reverse('maintenance_create', args=[panel.pk]),
        )

        self.assertContains(
            response,
            reverse('maintenance_detail', args=[signal_maintenance.pk]),
        )
        self.assertContains(response, 'Réparer cette face')
        self.assertContains(response, 'class="btn btn-sm btn-success"')
        self.assertNotContains(
            response,
            f'name="etat_face_{signal_face.pk}"',
        )
        self.assertContains(
            response,
            f'name="etat_face_{available_face.pk}"',
        )
        self.assertContains(response, '2. Rapport d\'Intervention')

        blocked_response = self.client.post(
            reverse('maintenance_create', args=[panel.pk]),
            {
                'support': panel.pk,
                'effectue_par': user.pk,
                'date_intervention': '2026-10-24T10:00',
                f'etat_face_{signal_face.pk}': Support.ETAT_BON,
                f'etat_face_{available_face.pk}': '',
            },
        )
        self.assertEqual(blocked_response.status_code, 200)
        self.assertEqual(
            Maintenance.objects.filter(support=panel).count(),
            1,
        )

        allowed_response = self.client.post(
            reverse('maintenance_create', args=[panel.pk]),
            {
                'support': panel.pk,
                'effectue_par': user.pk,
                'date_intervention': '2026-10-24T10:00',
                f'etat_face_{signal_face.pk}': '',
                f'etat_face_{available_face.pk}': Support.ETAT_BON,
            },
        )
        self.assertEqual(allowed_response.status_code, 302)
        self.assertEqual(
            Maintenance.objects.filter(support=panel).count(),
            2,
        )
        available_face.refresh_from_db()
        self.assertEqual(available_face.etat, Support.ETAT_BON)

        Maintenance.objects.create(
            support=panel,
            face=available_face,
            etat_apres=Support.ETAT_A_MAINTENIR,
        )
        all_blocked_response = self.client.get(
            reverse('maintenance_create', args=[panel.pk]),
        )
        self.assertTrue(all_blocked_response.context['all_faces_blocked'])
        self.assertContains(
            all_blocked_response,
            'id="bloc-rapport" class="card shadow-sm mb-4" hidden',
            html=False,
        )
        self.assertContains(
            all_blocked_response,
            'id="bloc-preuve" class="card shadow-sm mb-4" hidden',
            html=False,
        )
        self.assertContains(
            all_blocked_response,
            'id="submit-footer" class="card-footer bg-light py-3 px-4 d-flex align-items-center justify-content-between flex-wrap gap-2 border-top" hidden',
            html=False,
        )
        self.assertContains(
            all_blocked_response,
            reverse('maintenance_detail', args=[signal_maintenance.pk]),
        )
        self.assertContains(
            all_blocked_response,
            reverse(
                'maintenance_detail',
                args=[Maintenance.objects.filter(face=available_face).latest('pk').pk],
            ),
        )

    def test_blocked_support_cannot_receive_new_maintenance_until_repaired(self):
        user = get_user_model().objects.create_user(
            username='maintenance-blocked-create',
            password='test-password',
            role='admin',
        )
        self.client.force_login(user)

        for index, state in enumerate(
            (Support.ETAT_A_MAINTENIR, Support.ETAT_PANNE),
            start=1,
        ):
            with self.subTest(state=state):
                screen = Support.objects.create(
                    code=f'MAINT-BLOCKED-{index}',
                    nom='Écran à réparer',
                    type_support=Support.TYPE_ECRAN,
                    etat=state,
                    latitude=12.34,
                    longitude=-1.52,
                )
                existing = Maintenance.objects.create(
                    support=screen,
                    etat_apres=state,
                )
                initial_count = Maintenance.objects.filter(
                    support=screen,
                ).count()

                response = self.client.get(
                    reverse('maintenance_create', args=[screen.pk]),
                )
                self.assertContains(response, screen.get_etat_display())
                self.assertContains(response, 'Réparez ce support')
                self.assertContains(response, 'disabled')

                response = self.client.post(
                    reverse('maintenance_create', args=[screen.pk]),
                    {
                        'support': screen.pk,
                        'effectue_par': user.pk,
                        'date_intervention': '2026-10-24T10:00',
                        'etat_apres': Support.ETAT_BON,
                    },
                )
                self.assertEqual(response.status_code, 200)
                self.assertEqual(
                    Maintenance.objects.filter(support=screen).count(),
                    initial_count,
                )
                self.assertContains(response, 'Réparez ce support')

                self.client.post(
                    reverse('maintenance_quick_repair', args=[existing.pk]),
                )
                screen.refresh_from_db()
                self.assertEqual(screen.etat, Support.ETAT_BON)

                response = self.client.post(
                    reverse('maintenance_create', args=[screen.pk]),
                    {
                        'support': screen.pk,
                        'effectue_par': user.pk,
                        'date_intervention': '2026-10-24T10:00',
                        'etat_apres': Support.ETAT_BON,
                    },
                )
                self.assertEqual(response.status_code, 302)

    def test_panel_state_tracks_preventive_and_failed_faces(self):
        panel = Support.objects.create(
            code='MAINT-PANEL',
            nom='Panneau à entretenir',
            type_support=Support.TYPE_PANNEAU,
            latitude=12.34,
            longitude=-1.52,
        )
        face_a = FacePanneau.objects.create(support=panel, label='A')
        face_b = FacePanneau.objects.create(support=panel, label='B')

        Maintenance.objects.create(
            support=panel,
            face=face_a,
            etat_apres=Support.ETAT_A_MAINTENIR,
        )
        panel.refresh_from_db()
        self.assertEqual(panel.etat, Support.ETAT_A_MAINTENIR)

        Maintenance.objects.create(
            support=panel,
            face=face_b,
            etat_apres=Support.ETAT_PANNE,
        )
        panel.refresh_from_db()
        self.assertEqual(panel.etat, Support.ETAT_A_MAINTENIR)

        Maintenance.objects.create(
            support=panel,
            face=face_a,
            etat_apres=Support.ETAT_PANNE,
        )
        panel.refresh_from_db()
        self.assertEqual(panel.etat, Support.ETAT_PANNE)

    def test_signal_state_reports_functional_face_needing_attention(self):
        panel = Support.objects.create(
            code='MAINT-PANEL-ACTIVE',
            nom='Panneau avec signal',
            type_support=Support.TYPE_PANNEAU,
            latitude=12.34,
            longitude=-1.52,
        )
        face_a = FacePanneau.objects.create(support=panel, label='A')
        FacePanneau.objects.create(support=panel, label='B')

        Maintenance.objects.create(
            support=panel,
            face=face_a,
            etat_apres=Support.ETAT_A_MAINTENIR,
        )

        panel.refresh_from_db()
        self.assertEqual(panel.etat, Support.ETAT_A_MAINTENIR)

    def test_signal_and_panne_form_one_period_until_bon_state(self):
        panel = Support.objects.create(
            code='MAINT-SIGNAL-PERIOD',
            nom='Panneau avec signal puis panne',
            type_support=Support.TYPE_PANNEAU,
            latitude=12.34,
            longitude=-1.52,
        )
        face = FacePanneau.objects.create(support=panel, label='A')
        signal_at = datetime.datetime(2025, 1, 1, 12, tzinfo=datetime.timezone.utc)
        panne_at = datetime.datetime(2025, 1, 2, 12, tzinfo=datetime.timezone.utc)
        bon_at = datetime.datetime(2025, 1, 3, 12, tzinfo=datetime.timezone.utc)
        Maintenance.objects.create(
            support=panel,
            face=face,
            etat_apres=Support.ETAT_A_MAINTENIR,
            date_intervention=signal_at,
        )
        Maintenance.objects.create(
            support=panel,
            face=face,
            etat_apres=Support.ETAT_PANNE,
            date_intervention=panne_at,
        )
        Maintenance.objects.create(
            support=panel,
            face=face,
            etat_apres=Support.ETAT_BON,
            date_intervention=bon_at,
        )

        periods = panel.get_periodes_panne()

        self.assertEqual(len(periods), 1)
        self.assertEqual(periods[0]['debut'], signal_at)
        self.assertEqual(periods[0]['fin'], bon_at)
        self.assertEqual(periods[0]['etat_debut'], Support.ETAT_A_MAINTENIR)
        self.assertEqual(periods[0]['etat_incident'], Support.ETAT_PANNE)
        self.assertEqual(
            face.jours_disponibles_sur_periode(
                datetime.date(2025, 1, 1),
                datetime.date(2025, 1, 3),
            ),
            1,
        )

    def test_global_periods_page_includes_open_signal(self):
        user = get_user_model().objects.create_user(
            username='maintenance-periods-admin',
            password='test-password',
            role='admin',
        )
        self.client.force_login(user)
        screen = Support.objects.create(
            code='MAINT-GLOBAL-SIGNAL',
            nom='Écran signalé',
            type_support=Support.TYPE_ECRAN,
            latitude=12.34,
            longitude=-1.52,
        )
        signal = Maintenance.objects.create(
            support=screen,
            etat_apres=Support.ETAT_A_MAINTENIR,
            description='Signalement à vérifier',
        )

        response = self.client.get(reverse('periodes_panne_liste'))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context['nb_en_cours'], 1)
        period = response.context['periodes'][0]
        self.assertEqual(period['maint_debut'], signal)
        self.assertEqual(period['etat_debut'], Support.ETAT_A_MAINTENIR)
        self.assertEqual(period['etat_incident'], Support.ETAT_A_MAINTENIR)
        self.assertIsNone(period['fin'])
        self.assertContains(response, 'Début du signalement / de la panne')
        self.assertContains(response, 'SIGNAL')
        self.assertContains(response, 'Signalement à vérifier')

        response = self.client.get(
            reverse('support_periodes_panne', args=[screen.pk]),
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'SIGNAL')

    def test_global_periods_page_prioritizes_open_pannes_and_is_paginated(self):
        user = get_user_model().objects.create_user(
            username='maintenance-periods-pagination',
            password='test-password',
            role='admin',
        )
        self.client.force_login(user)

        for index in range(22):
            screen = Support.objects.create(
                code=f'PERIOD-SIGNAL-{index:02d}',
                nom='Écran avec signal à réparer',
                type_support=Support.TYPE_ECRAN,
                latitude=12.34,
                longitude=-1.52,
            )
            Maintenance.objects.create(
                support=screen,
                etat_apres=Support.ETAT_A_MAINTENIR,
            )

        panne_screens = []
        for index in range(2):
            screen = Support.objects.create(
                code=f'PERIOD-PANNE-{index:02d}',
                nom='Écran en panne',
                type_support=Support.TYPE_ECRAN,
                latitude=12.34,
                longitude=-1.52,
            )
            panne_screens.append(screen)
            Maintenance.objects.create(
                support=screen,
                etat_apres=Support.ETAT_PANNE,
            )

        response = self.client.get(
            reverse('periodes_panne_liste'),
            {'statut': 'en_cours'},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context['nb_total'], 24)
        self.assertEqual(response.context['page_obj'].paginator.num_pages, 2)
        self.assertCountEqual(
            [period['support'] for period in response.context['periodes'][:2]],
            panne_screens,
        )
        self.assertContains(response, 'Page 1 / 2')
        self.assertContains(response, 'statut=en_cours&amp;page=2')

    def test_period_report_filters_overlapping_periods_and_exports_pdf_excel(self):
        user = get_user_model().objects.create_user(
            username='maintenance-period-report',
            password='test-password',
            role='admin',
        )
        self.client.force_login(user)
        window_start = datetime.datetime(2025, 4, 10, tzinfo=datetime.timezone.utc)
        window_end = datetime.datetime(2025, 4, 12, 23, 59, tzinfo=datetime.timezone.utc)

        periods = [
            ('REPORT-OVERLAP-BEFORE', Support.ETAT_PANNE,
             datetime.datetime(2025, 4, 9, tzinfo=datetime.timezone.utc),
             datetime.datetime(2025, 4, 11, tzinfo=datetime.timezone.utc)),
            ('REPORT-OVERLAP-AFTER', Support.ETAT_A_MAINTENIR,
             datetime.datetime(2025, 4, 12, tzinfo=datetime.timezone.utc),
             datetime.datetime(2025, 4, 13, tzinfo=datetime.timezone.utc)),
            ('REPORT-OUTSIDE', Support.ETAT_PANNE,
             datetime.datetime(2025, 4, 13, tzinfo=datetime.timezone.utc),
             datetime.datetime(2025, 4, 14, tzinfo=datetime.timezone.utc)),
        ]
        for code, incident_state, incident_at, repaired_at in periods:
            screen = Support.objects.create(
                code=code,
                nom='Écran pour rapport',
                type_support=Support.TYPE_ECRAN,
                latitude=12.34,
                longitude=-1.52,
            )
            Maintenance.objects.create(
                support=screen,
                etat_apres=incident_state,
                date_intervention=incident_at,
                description='Constat initial',
            )
            Maintenance.objects.create(
                support=screen,
                etat_apres=Support.ETAT_BON,
                date_intervention=repaired_at,
                description='Réparation effectuée',
            )

        filters = {
            'date_debut': window_start.date().isoformat(),
            'date_fin': window_end.date().isoformat(),
        }
        response = self.client.get(reverse('periodes_panne_liste'), filters)

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context['nb_total'], 2)
        self.assertCountEqual(
            [period['support'].code for period in response.context['periodes']],
            ['REPORT-OVERLAP-BEFORE', 'REPORT-OVERLAP-AFTER'],
        )

        pdf_response = self.client.get(
            reverse('pannes_report_pdf'),
            filters,
        )
        self.assertEqual(pdf_response.status_code, 200)
        self.assertEqual(pdf_response['Content-Type'], 'application/pdf')
        self.assertTrue(pdf_response.content.startswith(b'%PDF'))

        excel_response = self.client.get(
            reverse('pannes_report_excel'),
            filters,
        )
        self.assertEqual(excel_response.status_code, 200)
        workbook = openpyxl.load_workbook(io.BytesIO(excel_response.content))
        self.assertEqual(workbook.active.max_row, 3)
        self.assertEqual(
            {workbook.active.cell(row=row, column=1).value for row in (2, 3)},
            {'REPORT-OVERLAP-BEFORE', 'REPORT-OVERLAP-AFTER'},
        )

        preview_response = self.client.get(
            reverse('pannes_report_preview'),
            filters,
        )
        self.assertEqual(preview_response.status_code, 200)
        self.assertEqual(preview_response.context['nb_total'], 2)
        self.assertContains(preview_response, 'Télécharger PDF')
        self.assertContains(preview_response, 'Excel')

    def test_legacy_maintenance_values_are_migrated_to_signal(self):
        panel = Support.objects.create(
            code='MAINT-LEGACY',
            nom='Panneau état historique',
            type_support=Support.TYPE_PANNEAU,
            latitude=12.34,
            longitude=-1.52,
        )
        face = FacePanneau.objects.create(support=panel, label='A')
        maintenance = Maintenance.objects.create(
            support=panel,
            face=face,
            etat_apres=Support.ETAT_BON,
        )
        Support.objects.filter(pk=panel.pk).update(etat='maintenance')
        FacePanneau.objects.filter(pk=face.pk).update(etat='maintenance')
        Maintenance.objects.filter(pk=maintenance.pk).update(
            etat_apres='maintenance',
        )

        migration = import_module(
            'inventory.migrations.0021_signal_support_state',
        )
        migration.migrate_maintenance_state_to_signal(apps, None)

        panel.refresh_from_db()
        face.refresh_from_db()
        maintenance.refresh_from_db()
        self.assertEqual(panel.etat, Support.ETAT_A_MAINTENIR)
        self.assertEqual(face.etat, Support.ETAT_A_MAINTENIR)
        self.assertEqual(maintenance.etat_apres, Support.ETAT_A_MAINTENIR)

    def test_preventive_state_remains_selectable_for_campaigns(self):
        screen = Support.objects.create(
            code='MAINT-BOOKABLE',
            nom='Écran réservable',
            type_support=Support.TYPE_ECRAN,
            etat=Support.ETAT_A_MAINTENIR,
            latitude=12.34,
            longitude=-1.52,
        )

        form = LigneCampagneForm(data={'type_support': Support.TYPE_ECRAN})

        self.assertIn(screen, form.fields['support'].queryset)


class SupportFormCategoryFilterTests(TestCase):
    def setUp(self):
        self.geant_format = FormatSupport.objects.create(
            code='FORM-GEANT',
            dimensions='8m × 6m',
            superficie=48,
            categorie='Géant',
        )
        self.sucette_format = FormatSupport.objects.create(
            code='FORM-SUCETTE',
            dimensions='1m × 2m',
            superficie=2,
            categorie='Sucette',
        )
        self.screen_format = FormatSupport.objects.create(
            code='FORM-ECRAN',
            dimensions='6 × 4',
            categorie='Ecran',
        )
        self.market_format = FormatSupport.objects.create(
            code='FORM-MARCHE',
            dimensions='Marché',
            categorie='Marché',
        )

    def test_form_lists_categories_and_groups_formats_by_category(self):
        form = SupportForm()
        choices = form.fields['format'].choices

        self.assertIn(('Géant', [('FORM-GEANT', '8m × 6m (48m²)')]), choices)
        self.assertIn(('Sucette', [('FORM-SUCETTE', '1m × 2m (2m²)')]), choices)
        self.assertIn(('Géant', 'Géant'), form.fields['categorie_format'].choices)
        self.assertIn(('Sucette', 'Sucette'), form.fields['categorie_format'].choices)
        self.assertIn(('Marché', 'Marché'), form.fields['categorie_format'].choices)
        self.assertNotIn(('Ecran', 'Ecran'), form.fields['categorie_format'].choices)
        self.assertNotIn(
            ('Ecran', [('FORM-ECRAN', '6 × 4')]),
            choices,
        )

    def test_selected_category_must_match_selected_format(self):
        form = SupportForm(data={
            'categorie_format': 'Géant',
            'format': self.sucette_format.code,
            'type_support': Support.TYPE_PANNEAU,
            'nb_faces': 2,
        })

        self.assertIn('format', form.errors)
        self.assertIn('ne correspond pas à la catégorie', form.errors['format'][0])


class SupportMarketPlacementDefaultsTests(TestCase):
    def setUp(self):
        self.market = Marche.objects.create(nom='Marché sans emplacement')
        self.market_format = FormatSupport.objects.create(
            code='MARKET-PANEL-FORMAT',
            dimensions='Emplacement marché',
            superficie=1,
            categorie='Marché',
        )
        staff = get_user_model().objects.create_user(
            username='market-panel-create-staff',
            password='test-password',
            role='staff',
        )
        self.client.force_login(staff)

    def test_new_support_form_exposes_market_and_editable_market_face_default(self):
        response = self.client.get(reverse('support_create'))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'id="id_marche_select" name="marche"')
        self.assertContains(response, "nbFacesInput.value = '1'")
        self.assertContains(response, 'marketFaceDefaultApplied = true;')
        self.assertContains(response, 'updateFaces();')
        self.assertContains(response, 'premierEmplacementLibre')
        self.assertContains(response, 'supportSubmitButton.disabled = blocked;')
        self.assertContains(response, 'Ce marché n’a aucun emplacement libre.')

    def test_cannot_create_panel_in_market_without_free_location(self):
        response = self.client.post(reverse('support_create'), {
            'type_choice': Support.TYPE_PANNEAU,
            'type_support': Support.TYPE_PANNEAU,
            'code': 'MARKET-PANEL-NO-SLOT',
            'nom': 'Panneau marché sans emplacement',
            'format': self.market_format.code,
            'categorie_format': 'Marché',
            'nb_faces': 1,
            'latitude': '12.3400000',
            'longitude': '-1.5200000',
            'ville': 'Ouagadougou',
            'etat': Support.ETAT_BON,
            'marche': str(self.market.pk),
            'emplacement': '',
        })

        self.assertEqual(response.status_code, 200)
        self.assertIn(
            "Ce marché n'a plus d'emplacement libre",
            response.context['form'].non_field_errors()[0],
        )
        self.assertFalse(
            Support.objects.filter(code='MARKET-PANEL-NO-SLOT').exists()
        )

    def test_market_emplacements_include_availability_in_code_order(self):
        market = Marche.objects.create(nom='Marché des emplacements')
        occupied = Emplacement.objects.create(marche=market, code='A-01')
        first_free = Emplacement.objects.create(marche=market, code='A-02')
        Emplacement.objects.create(marche=market, code='A-03')
        Support.objects.create(
            code='MARKET-OCCUPIED',
            nom='Support déjà installé',
            type_support=Support.TYPE_PANNEAU,
            emplacement=occupied,
            latitude=12.34,
            longitude=-1.52,
        )
        response = self.client.get(
            reverse('api_emplacements_marche', args=[market.pk])
        )

        self.assertEqual(response.status_code, 200)
        emplacements = response.json()['emplacements']
        self.assertEqual(
            [emplacement['code'] for emplacement in emplacements],
            ['A-01', 'A-02', 'A-03'],
        )
        self.assertFalse(emplacements[0]['libre'])
        self.assertEqual(
            next(emplacement['id'] for emplacement in emplacements if emplacement['libre']),
            first_free.pk,
        )

    def test_market_panel_is_included_in_supports_geojson(self):
        market = Marche.objects.create(
            nom='Marché visible sur la carte',
            latitude=12.345,
            longitude=-1.523,
        )
        emplacement = Emplacement.objects.create(marche=market, code='B-01')
        support = Support.objects.create(
            code='MARKET-PANEL-GEOJSON',
            nom='Panneau du marché',
            type_support=Support.TYPE_PANNEAU,
            format=self.market_format.code,
            emplacement=emplacement,
            latitude=12.345,
            longitude=-1.523,
        )
        second_emplacement = Emplacement.objects.create(marche=market, code='B-02')
        second_support = Support.objects.create(
            code='MARKET-PANEL-GEOJSON-2',
            nom='Deuxième panneau du marché',
            type_support=Support.TYPE_PANNEAU,
            format=self.market_format.code,
            emplacement=second_emplacement,
            latitude=0,
            longitude=0,
        )

        response = self.client.get(reverse('api_geojson'))

        self.assertEqual(response.status_code, 200)
        features = {
            feature['properties']['id']: feature
            for feature in response.json()['features']
        }
        feature = features[support.pk]
        second_feature = features[second_support.pk]
        self.assertEqual(feature['geometry']['coordinates'], [-1.523, 12.345])
        self.assertEqual(feature['properties']['marche'], market.nom)
        self.assertEqual(feature['properties']['marche_id'], market.pk)
        self.assertEqual(feature['properties']['emplacement_code'], emplacement.code)
        self.assertEqual(second_feature['geometry']['coordinates'], [-1.523, 12.345])
        self.assertEqual(second_feature['properties']['marche_id'], market.pk)


class FormatSupportDetailViewTests(TestCase):
    def test_detail_shows_format_details_and_matching_supports(self):
        format_support = FormatSupport.objects.create(
            code='DETAIL-4X3',
            dimensions='4m × 3m',
            superficie=12,
            categorie='Standard',
        )
        support = Support.objects.create(
            code='DETAIL-SUPPORT',
            nom='Panneau du format',
            type_support=Support.TYPE_PANNEAU,
            format=format_support.code,
            latitude=12.34,
            longitude=-1.52,
            ville='Ouagadougou',
            quartier='Centre',
        )

        request = RequestFactory().get(reverse(
            'format_support_detail', args=[format_support.pk]
        ))
        view = FormatSupportDetailView()
        view.setup(request, pk=format_support.pk)
        view.object = format_support
        context = view.get_context_data(object=format_support)
        rendered = render_to_string(view.template_name, context, request=request)

        self.assertEqual(context['support_count'], 1)
        self.assertEqual(list(context['supports']), [support])
        self.assertIn('4m × 3m', rendered)
        self.assertIn('Panneau du format', rendered)

    def test_detail_hides_empty_category(self):
        format_support = FormatSupport.objects.create(
            code='DETAIL-NO-CATEGORY',
            dimensions='4m × 3m',
            categorie='',
        )
        request = RequestFactory().get(reverse(
            'format_support_detail', args=[format_support.pk]
        ))
        request.user = AnonymousUser()
        view = FormatSupportDetailView()
        view.setup(request, pk=format_support.pk)
        view.object = format_support
        context = view.get_context_data(object=format_support)
        rendered = render_to_string(view.template_name, context, request=request)

        self.assertNotIn('Sans catégorie', rendered)
        self.assertNotIn('bi-tag-fill', rendered)
        self.assertNotIn('>Catégorie</dt>', rendered)

    def test_format_used_by_panel_cannot_be_deleted(self):
        format_support = FormatSupport.objects.create(
            code='DELETE-PROTECTED',
            dimensions='4m × 3m',
            superficie=12,
            categorie='Standard',
        )
        Support.objects.create(
            code='DELETE-PROTECTED-PANEL',
            nom='Panneau utilisant le format',
            type_support=Support.TYPE_PANNEAU,
            format=format_support.code,
            latitude=12.34,
            longitude=-1.52,
        )
        user = get_user_model().objects.create_user(
            username='format-delete-test',
            password='test-password',
            role='staff',
        )
        self.client.force_login(user)

        response = self.client.post(reverse('format_delete', args=[format_support.pk]))

        self.assertRedirects(
            response,
            reverse('format_support_detail', args=[format_support.pk]),
        )
        self.assertTrue(FormatSupport.objects.filter(pk=format_support.pk).exists())
        self.assertIn(
            'au moins un panneau l’utilise',
            str(list(get_messages(response.wsgi_request))),
        )

    def test_unused_format_can_still_be_deleted(self):
        format_support = FormatSupport.objects.create(
            code='DELETE-UNUSED',
            dimensions='4m × 3m',
            superficie=12,
            categorie='Standard',
        )
        user = get_user_model().objects.create_user(
            username='format-delete-unused-test',
            password='test-password',
            role='staff',
        )
        self.client.force_login(user)

        response = self.client.post(reverse('format_delete', args=[format_support.pk]))

        self.assertRedirects(response, reverse('format_list'))
        self.assertFalse(FormatSupport.objects.filter(pk=format_support.pk).exists())

    def test_detail_supports_are_paginated(self):
        format_support = FormatSupport.objects.create(
            code='DETAIL-PAGED',
            dimensions='4m × 3m',
            categorie='Standard',
        )
        for number in range(21):
            Support.objects.create(
                code=f'DETAIL-PAGED-{number:02d}',
                nom=f'Panneau {number:02d}',
                type_support=Support.TYPE_PANNEAU,
                format=format_support.code,
                latitude=12.34,
                longitude=-1.52,
                ville='Ouagadougou',
            )

        request = RequestFactory().get(
            reverse('format_support_detail', args=[format_support.pk]),
            {'page': '2'},
        )
        view = FormatSupportDetailView()
        view.setup(request, pk=format_support.pk)
        view.object = format_support
        context = view.get_context_data(object=format_support)
        rendered = render_to_string(view.template_name, context, request=request)

        self.assertEqual(context['paginator'].count, 21)
        self.assertEqual(context['page_obj'].number, 2)
        self.assertEqual(len(context['supports']), 1)
        self.assertIn('Page 2 / 2', rendered)
        self.assertIn('?page=1', rendered)

    def test_format_list_is_paginated_and_reports_total_count(self):
        for number in range(21):
            FormatSupport.objects.create(
                code=f'LIST-PAGED-{number:02d}',
                dimensions='4m × 3m',
                categorie='Standard',
            )

        request = RequestFactory().get(reverse('format_list'), {'page': '2'})
        view = FormatSupportListView()
        view.setup(request)
        object_list = view.get_queryset()
        context = view.get_context_data(object_list=object_list)
        rendered = render_to_string(view.template_name, context, request=request)

        self.assertEqual(context['paginator'].count, 21)
        self.assertEqual(context['page_obj'].number, 2)
        self.assertIn('21 format', rendered)
        self.assertIn('Page 2 / 2', rendered)

    def test_format_list_is_sorted_by_category_then_measure_and_code(self):
        formats = [
            ('SORT-LARGE', '8m × 6m', 48, 'Géant'),
            ('SORT-SMALL', '1m × 2m', 2, 'Sucette'),
            ('SORT-MIDDLE', '4m × 3m', 12, 'Standard'),
            ('SORT-DERIVED', '9m × 5m', None, 'Géant'),
            ('SORT-UNKNOWN', 'Personnalisé', None, 'Marché'),
        ]
        for code, dimensions, superficie, categorie in formats:
            FormatSupport.objects.create(
                code=code,
                dimensions=dimensions,
                superficie=superficie,
                categorie=categorie,
            )

        view = FormatSupportListView()
        view.setup(RequestFactory().get(
            reverse('format_list'),
            {'sort': 'code', 'dir': 'desc'},
        ))

        self.assertEqual(
            [format_support.code for format_support in view.get_queryset()],
            [
                'SORT-DERIVED',
                'SORT-LARGE',
                'SORT-UNKNOWN',
                'SORT-MIDDLE',
                'SORT-SMALL',
            ],
        )

    def test_format_list_defaults_to_category_then_superficie_order(self):
        expected = [
            ('ORDER-GEANT-SMALL', 2, 'Géant'),
            ('ORDER-GEANT-LARGE', 48, 'Géant'),
            ('ORDER-STANDARD', 12, 'Standard'),
            ('ORDER-SUCETTE', 1, 'Sucette'),
        ]
        for code, superficie, categorie in expected:
            FormatSupport.objects.create(
                code=code,
                dimensions=f'{superficie}m²',
                superficie=superficie,
                categorie=categorie,
            )

        view = FormatSupportListView()
        view.setup(RequestFactory().get(reverse('format_list')))

        ordered_formats = [
            item for item in view.get_queryset()
            if item.code.startswith('ORDER-')
        ]
        self.assertEqual(
            [(item.code, item.superficie) for item in ordered_formats],
            [(code, superficie) for code, superficie, _ in expected],
        )


class SupportDeletionApprovalTests(TestCase):
    def setUp(self):
        self.support = Support.objects.create(
            code='DELETE-APPROVAL',
            nom='Support à supprimer',
            type_support=Support.TYPE_ECRAN,
            latitude=12.34,
            longitude=-1.52,
        )
        self.requester = get_user_model().objects.create_user(
            username='deletion-requester',
            password='test-password',
            role='staff',
        )
        self.reviewer = get_user_model().objects.create_user(
            username='deletion-reviewer',
            password='test-password',
            role='staff',
        )
        self.admin = get_user_model().objects.create_user(
            username='deletion-admin',
            password='test-password',
            role='admin',
        )
        self.url = reverse('support_delete', args=[self.support.uuid])

    def test_second_staff_user_must_approve_and_cannot_be_requester(self):
        self.client.force_login(self.requester)

        confirmation_response = self.client.get(self.url)
        self.assertContains(confirmation_response, 'un autre compte utilisateur')

        first_response = self.client.post(self.url)

        self.assertRedirects(
            first_response,
            reverse('support_detail', args=[self.support.uuid]),
        )
        deletion_request = SupportDeletionRequest.objects.get(support=self.support)
        self.assertEqual(deletion_request.requested_by, self.requester)
        cache.clear()
        reviewer_request = RequestFactory().get('/')
        reviewer_request.user = self.reviewer
        reviewer_alerts = alerts_data(reviewer_request)['alertes']
        self.assertTrue(any(self.support.code in alert['msg'] for alert in reviewer_alerts))
        requester_request = RequestFactory().get('/')
        requester_request.user = self.requester
        requester_alerts = alerts_data(requester_request)['alertes']
        self.assertFalse(any(self.support.code in alert['msg'] for alert in requester_alerts))

        self.client.post(self.url)
        self.assertTrue(Support.objects.filter(pk=self.support.pk).exists())
        self.assertEqual(
            SupportDeletionRequest.objects.get(pk=deletion_request.pk).status,
            SupportDeletionRequest.STATUS_PENDING,
        )

        self.client.force_login(self.reviewer)
        queue_response = self.client.get(reverse('support_deletion_requests'))
        self.assertContains(queue_response, self.support.code)
        support_list_response = self.client.get(reverse('support_list'))
        self.assertContains(support_list_response, '1 demande à approuver')
        self.assertContains(support_list_response, 'Suppression en attente')
        approval_response = self.client.post(self.url)

        self.assertRedirects(approval_response, reverse('support_list'))
        self.assertFalse(Support.objects.filter(pk=self.support.pk).exists())
        deletion_request.refresh_from_db()
        self.assertIsNone(deletion_request.support)
        self.assertEqual(deletion_request.status, SupportDeletionRequest.STATUS_APPROVED)
        self.assertEqual(deletion_request.approved_by, self.reviewer)

    def test_admin_can_delete_support_without_second_approval(self):
        self.client.force_login(self.admin)

        response = self.client.post(self.url)

        self.assertRedirects(response, reverse('support_list'))
        self.assertFalse(Support.objects.filter(pk=self.support.pk).exists())
        self.assertFalse(SupportDeletionRequest.objects.filter(support_code=self.support.code).exists())

    def test_panel_also_requires_two_distinct_staff_users(self):
        panel = Support.objects.create(
            code='DELETE-PANEL-APPROVAL',
            nom='Panneau à supprimer',
            type_support=Support.TYPE_PANNEAU,
            format='4x3',
            latitude=12.34,
            longitude=-1.52,
        )
        panel_url = reverse('support_delete', args=[panel.uuid])
        self.client.force_login(self.requester)
        self.client.post(panel_url)
        self.client.force_login(self.reviewer)

        response = self.client.post(panel_url)

        self.assertEqual(response.status_code, 302)
        self.assertEqual(response.url, reverse('support_list'))
        self.assertFalse(Support.objects.filter(pk=panel.pk).exists())

    def test_support_linked_to_campaign_cannot_be_deleted_even_by_admin(self):
        client = Client.objects.create(
            nom='Client suppression',
            telephone='00000000',
            email='delete-campaign@example.com',
        )
        campaign = Campagne.objects.create(
            client=client,
            nom='Campagne liée',
            date_debut=datetime.date.today(),
            date_fin=datetime.date.today() + datetime.timedelta(days=1),
        )
        LigneCampagne.objects.create(campagne=campaign, support=self.support)
        self.client.force_login(self.admin)

        response = self.client.post(self.url)

        self.assertRedirects(
            response,
            reverse('support_detail', args=[self.support.uuid]),
        )
        self.assertTrue(Support.objects.filter(pk=self.support.pk).exists())
        self.assertFalse(SupportDeletionRequest.objects.filter(support=self.support).exists())

    def test_admin_can_delete_directly_and_cancels_pending_request(self):
        self.client.force_login(self.requester)
        self.client.post(self.url)
        self.client.force_login(self.admin)

        response = self.client.post(self.url)

        self.assertRedirects(response, reverse('support_list'))
        self.assertFalse(Support.objects.filter(pk=self.support.pk).exists())
        deletion_request = SupportDeletionRequest.objects.get(support_code=self.support.code)
        self.assertEqual(deletion_request.status, SupportDeletionRequest.STATUS_CANCELLED)


class SupportListCategoryFilterTests(TestCase):
    def setUp(self):
        self.geant_format = FormatSupport.objects.create(
            code='FILTER-GEANT',
            dimensions='8m × 6m',
            categorie='Géant',
        )
        self.sucette_format = FormatSupport.objects.create(
            code='FILTER-SUCETTE',
            dimensions='1m × 2m',
            categorie='Sucette',
        )
        self.screen_format = FormatSupport.objects.create(
            code='FILTER-ECRAN',
            dimensions='6 × 4',
            categorie='Ecran',
        )
        self.geant = Support.objects.create(
            code='FILTER-PANEL-GEANT',
            nom='Panneau géant',
            type_support=Support.TYPE_PANNEAU,
            format=self.geant_format.code,
            latitude=12.34,
            longitude=-1.52,
        )
        self.sucette = Support.objects.create(
            code='FILTER-PANEL-SUCETTE',
            nom='Panneau sucette',
            type_support=Support.TYPE_PANNEAU,
            format=self.sucette_format.code,
            latitude=12.34,
            longitude=-1.52,
            ville='Bobo-Dioulasso',
        )
        self.screen = Support.objects.create(
            code='FILTER-SCREEN',
            nom='Écran numérique',
            type_support=Support.TYPE_ECRAN,
            latitude=12.34,
            longitude=-1.52,
        )
        EcranNumerique.objects.create(
            support=self.screen,
            cellule=self.screen_format.code,
        )

    def get_view(self, query):
        request = RequestFactory().get(reverse('support_list'), query)
        view = SupportListView()
        view.setup(request)
        return view

    def test_category_filter_returns_supports_of_selected_category(self):
        view = self.get_view({'categorie': 'Géant'})

        self.assertEqual(list(view.get_queryset()), [self.geant])

    def test_12m2_menu_category_matches_standard_formats(self):
        standard_format = FormatSupport.objects.create(
            code='FILTER-STANDARD',
            dimensions='4m × 3m',
            superficie=12,
            categorie='Standard',
        )
        standard = Support.objects.create(
            code='FILTER-PANEL-STANDARD',
            nom='Panneau standard',
            type_support=Support.TYPE_PANNEAU,
            format=standard_format.code,
            latitude=12.34,
            longitude=-1.52,
        )

        view = self.get_view({'categorie': '12m2'})

        self.assertEqual(list(view.get_queryset()), [standard])

    def test_screen_category_filters_by_screen_cell_format(self):
        legacy_screen = Support.objects.create(
            code='FILTER-SCREEN-LEGACY',
            nom='Écran sans fiche technique',
            type_support=Support.TYPE_ECRAN,
            latitude=12.34,
            longitude=-1.52,
        )
        view = self.get_view({'categorie': 'Ecran'})

        self.assertEqual(set(view.get_queryset()), {self.screen, legacy_screen})

    def test_unknown_screen_cell_format_remains_filterable(self):
        screen = Support.objects.create(
            code='FILTER-SCREEN-UNKNOWN',
            nom='Écran format non catalogué',
            type_support=Support.TYPE_ECRAN,
            latitude=12.34,
            longitude=-1.52,
        )
        EcranNumerique.objects.create(support=screen, cellule='9x5')

        view = self.get_view({
            'categorie': 'Ecran',
            'type_panneau': '9x5',
        })
        context = view.get_context_data(object_list=view.get_queryset())

        self.assertEqual(list(view.get_queryset()), [screen])
        self.assertIn(('9x5', '9x5'), context['type_panneau_choices'])

    def test_categories_and_formats_are_limited_to_selected_city(self):
        view = self.get_view({
            'ville': 'Bobo-Dioulasso',
            'categorie': 'Sucette',
        })
        context = view.get_context_data(object_list=view.get_queryset())

        self.assertEqual(context['category_choices'], ['Sucette'])
        self.assertEqual(
            context['type_panneau_choices'],
            [('FILTER-SUCETTE', '1m × 2m')],
        )
        self.assertEqual(list(view.get_queryset()), [self.sucette])

    def test_category_choices_and_formats_are_exposed_to_filter(self):
        view = self.get_view({'categorie': 'Géant'})
        queryset = view.get_queryset()
        context = view.get_context_data(object_list=queryset)
        view.request.user = AnonymousUser()
        rendered = render_to_string(
            view.template_name,
            context,
            request=view.request,
        )

        self.assertIn('Géant', context['category_choices'])
        self.assertEqual(context['category_f'], 'Géant')
        self.assertEqual(
            context['type_panneau_choices'],
            [('FILTER-GEANT', '8m × 6m')],
        )
        self.assertIn('name="categorie"', rendered)
        self.assertIn('<option value="Géant" selected>Géant</option>', rendered)
        self.assertNotIn('Tous les types', rendered)


class InventoryNavigationTests(TestCase):
    def test_menu_keeps_standard_and_giant_by_city_and_globalizes_sucette_screen(self):
        geant_format = FormatSupport.objects.create(
            code='MENU-GEANT',
            dimensions='8m × 6m',
            superficie=48,
            categorie='Géant',
        )
        standard_format = FormatSupport.objects.create(
            code='MENU-STANDARD',
            dimensions='4m × 3m',
            superficie=12,
            categorie='Standard',
        )
        sucette_format = FormatSupport.objects.create(
            code='MENU-SUCETTE',
            dimensions='1m × 2m',
            superficie=2,
            categorie='Sucette',
        )
        Support.objects.create(
            code='MENU-PANEL-GEANT',
            nom='Panneau géant',
            type_support=Support.TYPE_PANNEAU,
            format=geant_format.code,
            latitude=12.34,
            longitude=-1.52,
            ville='Ouagadougou',
        )
        Support.objects.create(
            code='MENU-PANEL-STANDARD',
            nom='Panneau standard',
            type_support=Support.TYPE_PANNEAU,
            format=standard_format.code,
            latitude=12.34,
            longitude=-1.52,
            ville='Ouagadougou',
        )
        Support.objects.create(
            code='MENU-PANEL-SUCETTE',
            nom='Panneau sucette',
            type_support=Support.TYPE_PANNEAU,
            format=sucette_format.code,
            latitude=12.34,
            longitude=-1.52,
            ville='Bobo-Dioulasso',
        )
        for number in range(3):
            Support.objects.create(
                code=f'MENU-PANEL-BOBO-{number}',
                nom=f'Panneau géant {number}',
                type_support=Support.TYPE_PANNEAU,
                format=geant_format.code,
                latitude=12.34,
                longitude=-1.52,
                ville='Bobo-Dioulasso',
            )
        screen_format = FormatSupport.objects.create(
            code='MENU-ECRAN',
            dimensions='6 × 4',
            categorie='Ecran',
        )
        smaller_screen_format = FormatSupport.objects.create(
            code='MENU-ECRAN-2',
            dimensions='2 × 1',
            superficie=2,
            categorie='Ecran',
        )
        screen = Support.objects.create(
            code='MENU-SCREEN',
            nom='Écran de Ouagadougou',
            type_support=Support.TYPE_ECRAN,
            latitude=12.34,
            longitude=-1.52,
            ville='Ouagadougou',
        )
        EcranNumerique.objects.create(support=screen, cellule=screen_format.code)
        smaller_screen = Support.objects.create(
            code='MENU-SCREEN-2',
            nom='Petit écran de Ouagadougou',
            type_support=Support.TYPE_ECRAN,
            latitude=12.34,
            longitude=-1.52,
            ville='Ouagadougou',
        )
        EcranNumerique.objects.create(
            support=smaller_screen,
            cellule=smaller_screen_format.code,
        )

        navigation = inventory_villes(RequestFactory().get('/'))
        self.assertEqual(
            [city['ville'] for city in navigation['inventory_categories']],
            ['Bobo-Dioulasso', 'Ouagadougou'],
        )
        self.assertEqual(
            [city['support_count'] for city in navigation['inventory_categories']],
            [3, 2],
        )
        ouagadougou = next(
            city for city in navigation['inventory_categories']
            if city['ville'] == 'Ouagadougou'
        )
        categories = {
            category['label']: category['formats']
            for category in ouagadougou['categories']
        }

        self.assertEqual(set(categories), {'12m2', 'Géant'})
        self.assertEqual(
            categories['12m2'],
            [{'code': 'MENU-STANDARD', 'label': '4m × 3m (12 m²)'}],
        )
        self.assertEqual(
            categories['Géant'],
            [{'code': 'MENU-GEANT', 'label': '8m × 6m (48 m²)'}],
        )
        global_categories = {
            category['label']: category['formats']
            for category in navigation['inventory_global_categories']
        }
        self.assertEqual(
            global_categories['Ecran'],
            [
                {'code': 'MENU-ECRAN-2', 'label': '2 × 1 (2 Cellules)'},
                {'code': 'MENU-ECRAN', 'label': '6 × 4'},
            ],
        )
        self.assertEqual(
            global_categories['Sucette'],
            [{'code': 'MENU-SUCETTE', 'label': '1m × 2m (2 m²)'}],
        )

        request = RequestFactory().get(
            reverse('support_list'),
            {'ville': 'Ouagadougou', 'categorie': 'Géant'},
        )
        request.user = get_user_model().objects.create_user(
            username='inventory-menu-test',
            password='test-password',
        )
        rendered_menu = render_to_string('base.html', {}, request=request)

        self.assertNotIn('Tous les supports de la ville', rendered_menu)
        self.assertIn(
            '?ville=Ouagadougou&categorie=G%C3%A9ant&type_panneau=MENU-GEANT&desactif=ville',
            rendered_menu,
        )
        self.assertIn(
            '?ville=Ouagadougou&categorie=12m2&type_panneau=MENU-STANDARD&desactif=ville',
            rendered_menu,
        )
        self.assertIn(
            '12m2 · 4m × 3m (12 m²)',
            rendered_menu,
        )
        self.assertNotIn(
            '<span class="nav-submenu-label">12m2</span>',
            rendered_menu,
        )
        self.assertIn(
            '?categorie=Ecran&type_panneau=MENU-ECRAN',
            rendered_menu,
        )
        self.assertNotIn(
            '?ville=Bobo-Dioulasso&categorie=G%C3%A9ant&type_panneau=MENU-SUCETTE',
            rendered_menu,
        )
        self.assertNotIn(
            '?ville=Ouagadougou&categorie=Ecran',
            rendered_menu,
        )


class EmplacementBulkCreateTests(TestCase):
    def setUp(self):
        self.marche = Marche.objects.create(nom='Grand Marché')
        self.user = get_user_model().objects.create_user(
            username='emplacement-bulk-staff',
            password='test-password',
            role='staff',
        )
        self.client.force_login(self.user)
        self.url = reverse('emplacement_create', args=[self.marche.pk])

    def test_creates_requested_count_using_market_sequence(self):
        Emplacement.objects.create(marche=self.marche, code='GM-04')

        response = self.client.post(self.url, {'nombre': 3, 'notes': 'Allée nord'})

        self.assertRedirects(response, reverse('marche_detail', args=[self.marche.pk]))
        created = list(self.marche.emplacements.order_by('code'))
        self.assertEqual(
            [emplacement.code for emplacement in created],
            ['GM-04', 'GM-05', 'GM-06', 'GM-07'],
        )
        self.assertEqual(
            {emplacement.notes for emplacement in created if emplacement.code != 'GM-04'},
            {'Allée nord'},
        )

    def test_market_name_generates_expected_code_prefix_and_preview(self):
        response = self.client.get(self.url)

        self.assertEqual(response.context['code_prefix'], 'GM')
        self.assertEqual(response.context['code_start'], 1)
        self.assertContains(response, 'data-prefix="GM"')
        self.assertContains(response, 'data-start="1"')

        other_market = Marche.objects.create(nom='Marché de 10')
        other_response = self.client.get(
            reverse('emplacement_create', args=[other_market.pk])
        )
        self.assertEqual(other_response.context['code_prefix'], 'M10')
        creation_response = self.client.post(
            reverse('emplacement_create', args=[other_market.pk]),
            {'nombre': 2},
        )
        self.assertRedirects(
            creation_response,
            reverse('marche_detail', args=[other_market.pk]),
        )
        self.assertEqual(
            list(other_market.emplacements.order_by('code').values_list('code', flat=True)),
            ['M10-01', 'M10-02'],
        )

    def test_next_code_endpoint_uses_market_prefix_and_sequence(self):
        Emplacement.objects.create(marche=self.marche, code='GM-08')

        response = self.client.get(
            reverse('emplacement_next_code', args=[self.marche.pk])
        )

        self.assertEqual(response.json(), {'code': 'GM-09'})

    def test_count_is_required_and_must_stay_within_bounds(self):
        for data in ({}, {'nombre': 0}, {'nombre': 501}, {'nombre': 'abc'}):
            with self.subTest(data=data):
                response = self.client.post(self.url, data)
                self.assertEqual(response.status_code, 200)
                self.assertContains(response, 'nombre')
                self.assertEqual(self.marche.emplacements.count(), 0)

    def test_code_collision_returns_explicit_error_without_partial_creation(self):
        Emplacement.objects.create(marche=self.marche, code='GM-01')
        with patch('inventory.views.generer_prochain_code_emplacement', return_value='GM-01'):
            response = self.client.post(self.url, {'nombre': 2})

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'codes suivants existent déjà')
        self.assertEqual(self.marche.emplacements.count(), 1)

    def test_non_staff_cannot_create_market_emplacements(self):
        client_user = get_user_model().objects.create_user(
            username='emplacement-bulk-technicien',
            password='test-password',
            role='technicien',
        )
        self.client.force_login(client_user)

        response = self.client.post(self.url, {'nombre': 1})

        self.assertEqual(response.status_code, 302)
        self.assertEqual(response.url, reverse('login'))
        self.assertEqual(self.marche.emplacements.count(), 0)


class EmplacementBulkActionTests(TestCase):
    def setUp(self):
        self.marche = Marche.objects.create(nom='Marché actions')
        self.other_marche = Marche.objects.create(nom='Autre marché')
        self.user = get_user_model().objects.create_user(
            username='emplacement-actions-staff',
            password='test-password',
            role='staff',
        )
        self.client.force_login(self.user)
        self.url = reverse('emplacement_bulk_action', args=[self.marche.pk])
        self.first = Emplacement.objects.create(
            marche=self.marche,
            code='MA-01',
            notes='Ancienne note',
        )
        self.second = Emplacement.objects.create(
            marche=self.marche,
            code='MA-02',
            notes='Autre note',
        )

    def test_bulk_edit_updates_notes_only_for_selected_market_emplacements(self):
        foreign = Emplacement.objects.create(
            marche=self.other_marche,
            code='AM-01',
            notes='Ne pas modifier',
        )

        response = self.client.post(self.url, {
            'action': 'edit_notes',
            'selected_ids': [self.first.pk, self.second.pk],
            'notes': 'Note commune',
        })

        self.assertRedirects(response, reverse('marche_detail', args=[self.marche.pk]))
        self.first.refresh_from_db()
        self.second.refresh_from_db()
        foreign.refresh_from_db()
        self.assertEqual(self.first.notes, 'Note commune')
        self.assertEqual(self.second.notes, 'Note commune')
        self.assertEqual(foreign.notes, 'Ne pas modifier')

    def test_market_detail_renders_select_all_and_per_row_checkboxes_inside_action_form(self):
        response = self.client.get(reverse('marche_detail', args=[self.marche.pk]))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'id="select-all-emplacements"')
        self.assertContains(response, f'value="{self.first.pk}"')
        self.assertContains(response, f'value="{self.second.pk}"')
        form_start = response.content.decode().index('id="emplacement-bulk-form"')
        checkbox = response.content.decode().index(
            f'value="{self.first.pk}"',
            form_start,
        )
        form_end = response.content.decode().index('</form>', form_start)
        self.assertLess(checkbox, form_end)

    def test_market_detail_paginates_emplacements_and_links_installed_panel(self):
        for number in range(3, 22):
            Emplacement.objects.create(
                marche=self.marche,
                code=f'MA-{number:02d}',
            )
        format_support = FormatSupport.objects.create(
            code='MARKET-DETAIL-FORMAT',
            dimensions='4m × 3m',
            superficie=12,
            categorie='Standard',
        )
        panel = Support.objects.create(
            code='MARKET-DETAIL-PANEL',
            nom='Panneau installé au marché',
            type_support=Support.TYPE_PANNEAU,
            format=format_support.code,
            latitude=12.34,
            longitude=-1.52,
            emplacement=self.first,
        )

        first_page = self.client.get(reverse('marche_detail', args=[self.marche.pk]))

        self.assertEqual(first_page.status_code, 200)
        self.assertEqual(first_page.context['paginator'].count, 21)
        self.assertEqual(len(first_page.context['emplacements_page']), 20)
        self.assertContains(first_page, 'Page 1 / 2')
        self.assertContains(first_page, reverse('support_detail', args=[panel.uuid]))
        self.assertContains(first_page, 'title="Voir le panneau associé"')

        second_page = self.client.get(
            reverse('marche_detail', args=[self.marche.pk]),
            {'page': 2},
        )
        self.assertEqual(second_page.status_code, 200)
        self.assertEqual(second_page.context['page_obj'].number, 2)
        self.assertEqual(len(second_page.context['emplacements_page']), 1)
        self.assertContains(second_page, 'Page 2 / 2')

    def test_bulk_delete_is_admin_only(self):
        self.client.raise_request_exception = False
        response = self.client.post(self.url, {
            'action': 'delete',
            'selected_ids': [self.first.pk],
        })

        self.assertEqual(response.status_code, 403)
        self.assertTrue(Emplacement.objects.filter(pk=self.first.pk).exists())

    def test_admin_can_bulk_delete_selected_emplacements(self):
        admin = get_user_model().objects.create_user(
            username='emplacement-actions-admin',
            password='test-password',
            role='admin',
        )
        self.client.force_login(admin)

        response = self.client.post(self.url, {
            'action': 'delete',
            'selected_ids': [self.first.pk],
        })

        self.assertRedirects(response, reverse('marche_detail', args=[self.marche.pk]))
        self.assertFalse(Emplacement.objects.filter(pk=self.first.pk).exists())
        self.assertTrue(Emplacement.objects.filter(pk=self.second.pk).exists())

    def test_bulk_delete_skips_emplacements_with_installed_support(self):
        admin = get_user_model().objects.create_user(
            username='emplacement-actions-admin-installed',
            password='test-password',
            role='admin',
        )
        self.client.force_login(admin)
        support = Support.objects.create(
            code='BULK-INSTALLED',
            nom='Panneau installé',
            type_support=Support.TYPE_PANNEAU,
            latitude=12.34,
            longitude=-1.52,
            ville='Ouagadougou',
            emplacement=self.first,
        )

        response = self.client.post(self.url, {
            'action': 'delete',
            'selected_ids': [self.first.pk, self.second.pk],
        }, follow=True)

        self.assertRedirects(response, reverse('marche_detail', args=[self.marche.pk]))
        self.assertTrue(Emplacement.objects.filter(pk=self.first.pk).exists())
        self.assertTrue(Support.objects.filter(pk=support.pk).exists())
        self.assertFalse(Emplacement.objects.filter(pk=self.second.pk).exists())
        self.assertContains(response, 'panneau installé')

    @patch('reports.views._build_context_emplacement_detail', return_value={})
    @patch('reports.views.HTML')
    def test_bulk_download_returns_zip_with_one_report_per_selected_item(
        self,
        html_mock,
        context_mock,
    ):
        html_mock.return_value.write_pdf.return_value = b'%PDF-test'

        response = self.client.post(self.url, {
            'action': 'download',
            'selected_ids': [self.first.pk, self.second.pk],
        })

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'application/zip')
        with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
            self.assertEqual(
                archive.namelist(),
                [
                    f'rapport_emplacement_ma-01_{self.first.pk}.pdf',
                    f'rapport_emplacement_ma-02_{self.second.pk}.pdf',
                ],
            )
            self.assertEqual(archive.read(archive.namelist()[0]), b'%PDF-test')
        self.assertEqual(context_mock.call_count, 2)
