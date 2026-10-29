import base64
import datetime
from datetime import date, timedelta, timezone as datetime_timezone
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import patch

from django.test import SimpleTestCase, TestCase
from django.template.loader import render_to_string
from django.urls import reverse
from openpyxl import load_workbook
from PIL import Image

from accounts.models import User
from campaigns.models import Campagne, Client, LigneCampagne, PreuveDiffusion
from inventory.models import Maintenance, Support
from piges.drive_piges import DrivePigesError
from piges.models import PigeFichier
from reports.pannes_report_view import _add_months
from reports.supports_report_view import _group_supports_by_quartier
from reports.views import (
    _build_context_campagne,
    _prepare_piges_drive_pdf,
    _prepare_visuels_pdf,
)


class PrepareVisuelsPdfTests(SimpleTestCase):
    def prepare_visuel(self, name, content):
        fichier = SimpleNamespace(
            name=name,
            url=f"/media/{name}",
            open=lambda mode: BytesIO(content),
        )
        campagne = SimpleNamespace(
            visuels=SimpleNamespace(
                all=lambda: [SimpleNamespace(fichier=fichier)],
            ),
        )
        return _prepare_visuels_pdf(campagne)[0]

    def test_large_jpeg_is_resized_and_compressed_for_pdf(self):
        original = BytesIO()
        Image.new("RGB", (3000, 2000), (120, 60, 80)).save(
            original,
            format="JPEG",
            quality=95,
        )

        result = self.prepare_visuel("photo.jpg", original.getvalue())
        optimized = base64.b64decode(result["src"].partition(",")[2])
        optimized_image = Image.open(BytesIO(optimized))

        self.assertEqual(optimized_image.size, (1200, 800))
        self.assertEqual(optimized_image.format, "JPEG")
        self.assertLess(len(optimized), len(original.getvalue()))

    def test_png_poster_keeps_lossless_format_and_is_resized(self):
        original = BytesIO()
        Image.new("RGB", (1800, 1200), (12, 90, 170)).save(original, format="PNG")

        result = self.prepare_visuel("affiche.png", original.getvalue())
        optimized = base64.b64decode(result["src"].partition(",")[2])
        optimized_image = Image.open(BytesIO(optimized))

        self.assertLessEqual(max(optimized_image.size), 1200)
        self.assertEqual(optimized_image.format, "PNG")


class CampagnePdfPreuvesTests(TestCase):
    def test_pdf_includes_download_link_to_campaign_drive_folder(self):
        client = Client.objects.create(
            nom="Client lien PDF",
            telephone="00000000",
            email="lien-pdf@example.com",
        )
        campagne = Campagne.objects.create(
            client=client,
            nom="Campagne lien PDF",
            date_debut=datetime.date(2026, 10, 1),
            date_fin=datetime.date(2026, 10, 31),
            lien_piges="https://drive.google.com/drive/folders/piges-pdf-test",
        )
        context = _build_context_campagne(campagne)
        context["preuves_diffusion_pdf"] = []
        context["piges_drive_pdf"] = []

        html = render_to_string("reports/campagne_pdf.html", context)

        self.assertIn(campagne.lien_piges, html)
        self.assertIn("Ouvrir le dossier Google Drive de la campagne", html)

    def test_contexte_pdf_inclut_les_preuves_de_la_mere_et_de_ses_enfants(self):
        client = Client.objects.create(
            nom="Client test PDF",
            telephone="00000000",
            email="pdf@example.com",
        )
        campagne = Campagne.objects.create(
            client=client,
            nom="Campagne mère",
            date_debut=datetime.date(2026, 10, 1),
            date_fin=datetime.date(2026, 10, 31),
            est_mere=True,
        )
        enfant = Campagne.objects.create(
            client=client,
            nom="Sous-campagne",
            date_debut=datetime.date(2026, 10, 1),
            date_fin=datetime.date(2026, 10, 31),
            campagne_parente=campagne,
        )
        for campagne_preuve in (campagne, enfant):
            ligne = LigneCampagne.objects.create(campagne=campagne_preuve)
            PreuveDiffusion.objects.create(
                campagne=campagne_preuve,
                ligne=ligne,
                photo="piges/preuve-test.jpg",
            )
            PigeFichier.objects.create(
                campagne=campagne_preuve,
                support_nom="Support photo",
                ville="Ouagadougou",
                date_prise=datetime.date(2026, 10, 17),
                nom_fichier="preuve.jpg",
                drive_file_id=f"drive-file-{campagne_preuve.pk}",
            )

        context = _build_context_campagne(campagne)

        self.assertEqual(
            set(context["preuves_diffusion"].values_list("campagne_id", flat=True)),
            {campagne.pk, enfant.pk},
        )
        self.assertEqual(
            set(context["piges_drive"].values_list("campagne_id", flat=True)),
            {campagne.pk, enfant.pk},
        )

    @patch("reports.views.download_photo")
    @patch("reports.views.get_drive_service")
    def test_drive_photos_are_embedded_in_pdf_as_optimized_data_urls(self, get_service, get_photo):
        client = Client.objects.create(
            nom="Client test PDF",
            telephone="00000000",
            email="photo-pdf@example.com",
        )
        campagne = Campagne.objects.create(
            client=client,
            nom="Campagne photo Drive",
            date_debut=datetime.date(2026, 10, 1),
            date_fin=datetime.date(2026, 10, 31),
        )
        pige = PigeFichier.objects.create(
            campagne=campagne,
            support_nom="Panneau Centre",
            ville="Ouagadougou",
            date_prise=datetime.date(2026, 10, 17),
            nom_fichier="preuve.jpg",
            drive_file_id="drive-file-pdf-embed-test",
        )
        original = BytesIO()
        Image.new("RGB", (2400, 1200), (100, 40, 80)).save(original, format="JPEG")
        get_photo.return_value = original.getvalue()

        prepared = _prepare_piges_drive_pdf([pige])

        self.assertEqual(get_service.call_count, 1)
        get_photo.assert_called_once_with(get_service.return_value, pige.drive_file_id)
        self.assertEqual(prepared[0]["pige"], pige)
        self.assertTrue(prepared[0]["src"].startswith("data:image/jpeg;base64,"))
        embedded = base64.b64decode(prepared[0]["src"].partition(",")[2])
        with Image.open(BytesIO(embedded)) as image:
            self.assertEqual(image.size, (1200, 600))

    @patch(
        "reports.views.get_drive_service",
        side_effect=DrivePigesError(
            "Le refresh token Google Drive a expiré ou a été révoqué."
        ),
    )
    def test_preview_shows_staff_error_when_drive_refresh_token_is_invalid(self, _get_service):
        client = Client.objects.create(
            nom="Client erreur OAuth",
            telephone="00000000",
            email="oauth-preview@example.com",
        )
        campagne = Campagne.objects.create(
            client=client,
            nom="Campagne erreur OAuth",
            date_debut=datetime.date(2026, 10, 1),
            date_fin=datetime.date(2026, 10, 31),
        )
        PigeFichier.objects.create(
            campagne=campagne,
            support_nom="Panneau Centre",
            ville="Ouagadougou",
            date_prise=datetime.date(2026, 10, 17),
            nom_fichier="preuve.jpg",
            drive_file_id="drive-file-invalid-refresh-test",
        )
        staff = User.objects.create_user(
            username="staff-oauth-preview",
            password="test-password",
            role=User.ROLE_STAFF,
        )
        self.client.force_login(staff)

        response = self.client.get(reverse("campagne_preview", args=[campagne.pk]))

        self.assertEqual(response.status_code, 503)
        self.assertContains(
            response,
            "Le refresh token Google Drive a expiré ou a été révoqué.",
            status_code=503,
        )
        self.assertContains(
            response,
            "Les preuves de diffusion enregistrées dans l'application",
            status_code=503,
        )


class GroupSupportsByQuartierTests(SimpleTestCase):
    def test_groups_supports_by_city_and_quartier_and_sorts_groups(self):
        supports = [
            {"ville": "Ouagadougou", "quartier": "Dapoya", "type": "Panneau", "nom": "Support B", "code": "B", "nb_faces": 2},
            {"ville": "Bobo-Dioulasso", "quartier": "Dapoya", "type": "Écran", "nom": "Écran A", "code": "A", "nb_faces": 1},
            {"ville": "Ouagadougou", "quartier": "Dapoya", "type": "Écran", "nom": "Écran A", "code": "A", "nb_faces": 1},
            {"ville": "", "quartier": "", "type": "Marché", "nom": "Marché sans zone", "code": "M", "nb_faces": 3},
        ]

        sections = _group_supports_by_quartier(supports)

        self.assertEqual(
            [section["nom"] for section in sections],
            [
                "Bobo-Dioulasso · Dapoya",
                "Ouagadougou · Dapoya",
                "Ville non renseignée · Quartier non renseigné",
            ],
        )
        self.assertEqual(sections[1]["total"], 2)
        self.assertEqual(sections[1]["unites"], 3)
        self.assertEqual(
            [item["code"] for item in sections[1]["items"]],
            ["B", "A"],
        )


class PannesDateFilterTests(TestCase):
    def setUp(self):
        user = User.objects.create_user(
            username="pannes-date-filter",
            password="test-password",
            role=User.ROLE_ADMIN,
        )
        self.client.force_login(user)

    def create_period(self, code, start, end=None):
        support = Support.objects.create(
            code=code,
            nom="Support de test",
            type_support=Support.TYPE_ECRAN,
            latitude=12.34,
            longitude=-1.52,
        )
        Maintenance.objects.create(
            support=support,
            etat_apres=Support.ETAT_PANNE,
            date_intervention=start,
        )
        if end:
            Maintenance.objects.create(
                support=support,
                etat_apres=Support.ETAT_BON,
                date_intervention=end,
            )
        return support

    def test_months_calculate_end_date_and_limit_the_report_range(self):
        self.create_period(
            "REPORT-IN-MONTH-RANGE",
            datetime.datetime(2026, 10, 31, 12, tzinfo=datetime_timezone.utc),
            datetime.datetime(2026, 11, 1, 12, tzinfo=datetime_timezone.utc),
        )
        self.create_period(
            "REPORT-AFTER-MONTH-RANGE",
            datetime.datetime(2026, 11, 2, 12, tzinfo=datetime_timezone.utc),
            datetime.datetime(2026, 11, 3, 12, tzinfo=datetime_timezone.utc),
        )

        response = self.client.get(
            reverse("pannes_report_preview"),
            {"date_debut": "2026-10-01", "duree_mois": "1"},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["date_fin_obj"], date(2026, 11, 1))
        self.assertEqual(response.context["filters"]["date_fin"], "2026-11-01")
        self.assertEqual(response.context["nb_total"], 1)
        self.assertIn("date_fin=2026-11-01", response.context["export_pdf_url"])

    def test_end_of_month_is_clamped_when_adding_months(self):
        self.assertEqual(_add_months(date(2024, 1, 31), 1), date(2024, 2, 29))

    def test_explicit_end_date_is_used_without_month_duration(self):
        response = self.client.get(
            reverse("pannes_report_preview"),
            {"date_debut": "2026-10-01", "date_fin": "2026-10-31"},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["date_fin_obj"], date(2026, 10, 31))
        self.assertContains(response, 'value="2026-10-31"')

    def test_month_duration_is_also_applied_to_the_periods_list(self):
        response = self.client.get(
            reverse("periodes_panne_liste"),
            {"date_debut": "2026-10-01", "duree_mois": "1"},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["filters"]["date_fin"], "2026-11-01")
        self.assertContains(response, 'value="2026-11-01"')
        self.assertIn("duree_mois=1", response.context["report_preview_url"])

    def test_open_period_does_not_match_a_future_date_range(self):
        from django.utils import timezone

        today = timezone.localdate()
        self.create_period(
            "REPORT-OPEN-BEFORE-FUTURE-RANGE",
            timezone.now() - timedelta(days=1),
        )

        response = self.client.get(
            reverse("pannes_report_preview"),
            {
                "date_debut": (today + timedelta(days=2)).isoformat(),
                "date_fin": (today + timedelta(days=3)).isoformat(),
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["nb_total"], 0)

    def test_report_distinguishes_unstarted_and_started_repairs(self):
        awaiting = self.create_period(
            "REPORT-AWAITING-REPAIR",
            datetime.datetime(2026, 10, 10, 12, tzinfo=datetime_timezone.utc),
        )
        started = self.create_period(
            "REPORT-STARTED-REPAIR",
            datetime.datetime(2026, 10, 11, 12, tzinfo=datetime_timezone.utc),
        )
        started_incident = Maintenance.objects.get(
            support=started,
            etat_apres=Support.ETAT_PANNE,
        )
        started_incident.reparation_commencee_le = datetime.datetime(
            2026,
            10,
            12,
            12,
            tzinfo=datetime_timezone.utc,
        )
        started_incident.save(update_fields=['reparation_commencee_le'])

        response = self.client.get(reverse("pannes_report_preview"))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["nb_a_reparer"], 1)
        self.assertEqual(response.context["nb_reparation_en_cours"], 1)
        periods_by_code = {
            period["support"].code: period
            for period in response.context["periodes"]
        }
        self.assertIsNone(
            periods_by_code[awaiting.code]["reparation_commencee_le"],
        )
        self.assertEqual(
            periods_by_code[started.code]["reparation_commencee_le"],
            started_incident.reparation_commencee_le,
        )
        self.assertContains(response, "En cours de réparation")
        self.assertContains(response, "À réparer")

        excel_response = self.client.get(reverse("pannes_report_excel"))
        self.assertEqual(excel_response.status_code, 200)
        workbook = load_workbook(
            BytesIO(excel_response.content),
            read_only=True,
        )
        export_rows = {
            row[0]: row
            for row in workbook.active.iter_rows(min_row=2, values_only=True)
        }
        self.assertEqual(export_rows[awaiting.code][7], "À réparer")
        self.assertEqual(
            export_rows[started.code][7],
            "En cours de réparation",
        )
        self.assertIsNotNone(export_rows[started.code][8])

        pdf_response = self.client.get(reverse("pannes_report_pdf"))
        self.assertEqual(pdf_response.status_code, 200)
        self.assertEqual(pdf_response["Content-Type"], "application/pdf")
        self.assertGreater(len(pdf_response.content), 0)